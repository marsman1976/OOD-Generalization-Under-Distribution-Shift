#!/usr/bin/env python3
"""
Phase 7.8 — Build the frozen plant-level Phase-7 analytical dataset

Creates one row per harvested lettuce plant and joins each plant to the
validated environmental history of its experiment × treatment × replicate
domain.

Frozen P7.7 rules implemented
------------------------------
- Experimental observation: individual plant.
- Expected plants: 216 (12 plants × 18 domains).
- Primary target: final total fresh plant weight (g).
- Primary domain predictors: pH, EC, TDS, water temperature history summaries.
- Portable measurements are summarized at domain level; hourly/portable rows
  are NOT treated as independent ML observations.
- EXP1 portable domain semantics use the documented P7.6b role swap:
      corrected_replicate = raw_portable_treatment
      corrected_treatment = raw_portable_replicate
- EXP2 raw portable dates A31/A32 are preserved in the archive but interpreted
  analytically as 13/05/2024 and 14/05/2024.
- No harvest-relative window is invented. Full validated observed portable
  experiment history is summarized.
- No synthetic oracle augmentation is applied.
- No model is trained here.

Outputs are written under:
results/experiment_7/plant_level_dataset/
"""

from __future__ import annotations

import hashlib
import json
import math
import re
import zipfile
from collections import defaultdict
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

import numpy as np
import openpyxl
import pandas as pd


ROOT = Path(__file__).resolve().parents[1]
ARCHIVE = ROOT / "data" / "phase7" / "raw" / "all_months_sensory_data.zip"
OUT = ROOT / "results" / "experiment_7" / "plant_level_dataset"

DATASET_CSV = OUT / "phase7_plant_level_dataset.csv"
ENV_CSV = OUT / "phase7_domain_environment_summaries.csv"
HARVEST_CSV = OUT / "phase7_harvest_plant_outcomes.csv"
PORTABLE_LONG_CSV = OUT / "phase7_portable_measurements_long.csv"
REPORT = OUT / "phase7_plant_level_dataset_report.txt"
MANIFEST = OUT / "phase7_plant_level_dataset_manifest.json"

EXPECTED_SHA256 = "daa6d1cc259234d4b1cb8ba9901b2f7e35975a444154d5b8dd0a58a394fdbf14"
EXPECTED_PLANTS = 216
EXPECTED_DOMAINS = 18
EXPECTED_PLANTS_PER_DOMAIN = 12

VARIABLES = ("ph", "ec", "tds", "water_temperature")
SUMMARY_STATS = ("mean", "std", "min", "max", "q25", "median", "q75", "slope")


def clean(v: Any) -> str:
    if v is None:
        return ""
    try:
        if pd.isna(v):
            return ""
    except Exception:
        pass
    return re.sub(r"\s+", " ", str(v)).strip()


def norm_token(v: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", clean(v).lower())


def to_float(v: Any) -> float:
    if v is None or isinstance(v, bool):
        return math.nan
    if isinstance(v, (int, float, np.number)):
        try:
            return float(v)
        except Exception:
            return math.nan
    s = clean(v).replace(",", ".")
    if not s or s in {"-", "—", "–"} or s.startswith("="):
        return math.nan
    m = re.fullmatch(r"[-+]?(?:\d+(?:\.\d*)?|\.\d+)", s)
    return float(s) if m else math.nan


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def experiment_number(name: str) -> int | None:
    m = re.search(r"EXP\.?\s*(\d+)", Path(name).name, flags=re.I)
    return int(m.group(1)) if m else None


def choose_sheet(names: list[str], hint: str) -> str | None:
    for n in names:
        if hint.lower() in n.lower():
            return n
    return None


def parse_domain_label(v: Any) -> tuple[int, int] | None:
    s = re.sub(r"[^A-Z0-9]+", "", clean(v).upper())
    m = re.fullmatch(r"R(\d+)T(\d+)", s)
    if m:
        return int(m.group(1)), int(m.group(2))
    return None


def parse_portable_label(v: Any) -> tuple[int, int] | None:
    s = clean(v).upper()
    r = re.search(r"\bREPLICATE\s*[-_:]?\s*(\d+)\b", s)
    t = re.search(r"\bT\s*[-_:]?\s*(\d+)\b", s)
    if r and t:
        return int(r.group(1)), int(t.group(1))
    return None


def corrected_domain(exp: int, raw_r: int, raw_t: int) -> tuple[int, int]:
    if exp == 1:
        return raw_t, raw_r
    return raw_r, raw_t


def parse_date_string_ddmmyyyy(s: str) -> pd.Timestamp | None:
    s = clean(s).replace("\\", "/").replace("-", "/")
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", s)
    if not m:
        return None
    d, mo, y = map(int, m.groups())
    try:
        return pd.Timestamp(year=y, month=mo, day=d)
    except Exception:
        return None


def analytical_portable_date(exp: int, cell) -> tuple[pd.Timestamp | None, str]:
    """
    Reconstruct portable dates using the source's DD/MM/YYYY convention.

    EXP2 A31/A32 receive the already-forensically-supported analytical
    interpretation. Raw workbook values remain untouched.

    For Excel datetime objects, month/day ambiguity is resolved by testing the
    ordinary and swapped interpretations and choosing continuity later at the
    sheet-sequence level.
    """
    if exp == 2 and cell.coordinate == "A31":
        return pd.Timestamp("2024-05-13"), "EXP2_A31_DOCUMENTED_ANALYTICAL_CORRECTION"
    if exp == 2 and cell.coordinate == "A32":
        return pd.Timestamp("2024-05-14"), "EXP2_A32_DOCUMENTED_ANALYTICAL_CORRECTION"

    v = cell.value
    if isinstance(v, datetime):
        return pd.Timestamp(v.date()), "EXCEL_DATETIME_INITIAL"
    if isinstance(v, date):
        return pd.Timestamp(v), "EXCEL_DATE_INITIAL"

    parsed = parse_date_string_ddmmyyyy(clean(v))
    if parsed is not None:
        return parsed, "EXPLICIT_DDMMYYYY"

    return None, "UNPARSED"


def repair_excel_date_sequence(records: list[dict]) -> None:
    """
    Correct Excel auto-converted date objects by choosing, for each ambiguous
    Excel date, between YYYY-MM-DD and YYYY-DD-MM using local chronology.

    Explicit DD/MM/YYYY strings and the documented EXP2 corrections are fixed
    anchors. This reproduces the P7.5b/P7.5c source convention without editing
    raw cells.
    """
    if not records:
        return

    # Candidate dates for each row.
    candidates: list[list[pd.Timestamp]] = []
    for rec in records:
        ts = rec["date"]
        prov = rec["date_provenance"]
        if ts is None:
            candidates.append([])
            continue
        if prov.startswith("EXCEL_") and ts.month <= 12 and ts.day <= 12 and ts.month != ts.day:
            normal = ts.normalize()
            try:
                swapped = pd.Timestamp(year=ts.year, month=ts.day, day=ts.month)
                candidates.append([normal, swapped])
            except Exception:
                candidates.append([normal])
        else:
            candidates.append([ts.normalize()])

    # Dynamic programming: prefer forward daily continuity and strongly punish
    # backward jumps. Dataset is tiny, so exhaustive two-state DP is simple.
    states: list[dict[int, tuple[float, int | None]]] = []
    for i, cand in enumerate(candidates):
        cur: dict[int, tuple[float, int | None]] = {}
        if not cand:
            states.append(cur)
            continue

        prev_i = i - 1
        while prev_i >= 0 and not states[prev_i]:
            prev_i -= 1

        if prev_i < 0:
            for j in range(len(cand)):
                cur[j] = (0.0, None)
        else:
            prev_cand = candidates[prev_i]
            for j, dt in enumerate(cand):
                best = None
                for k, (prev_cost, _) in states[prev_i].items():
                    delta = (dt - prev_cand[k]).days
                    if delta < 0:
                        penalty = 1000.0 + abs(delta) * 20.0
                    else:
                        penalty = abs(delta - (i - prev_i))
                    cost = prev_cost + penalty
                    if best is None or cost < best[0]:
                        best = (cost, k)
                cur[j] = best
        states.append(cur)

    last = len(records) - 1
    while last >= 0 and not states[last]:
        last -= 1
    if last < 0:
        return

    idx = min(states[last], key=lambda j: states[last][j][0])
    chosen: dict[int, int] = {}
    i = last
    while i >= 0:
        if states[i]:
            chosen[i] = idx
            prev_idx = states[i][idx][1]
            p = i - 1
            while p >= 0 and not states[p]:
                p -= 1
            if p < 0:
                break
            idx = prev_idx if prev_idx is not None else 0
            i = p
        else:
            i -= 1

    for i, j in chosen.items():
        old = records[i]["date"]
        new = candidates[i][j]
        if old is not None and new != old.normalize():
            records[i]["date"] = new
            records[i]["date_provenance"] += "|EXCEL_DAY_MONTH_SWAP"
        else:
            records[i]["date"] = new


def find_domain_for_column(ws, ph_row: int, col: int) -> tuple[int, int, str] | None:
    # Direct / merged labels above pH header.
    for hr in range(ph_row - 1, 0, -1):
        parsed = parse_portable_label(ws.cell(hr, col).value)
        if parsed:
            return parsed[0], parsed[1], ws.cell(hr, col).coordinate
        for rng in ws.merged_cells.ranges:
            if rng.min_row <= hr <= rng.max_row and rng.min_col <= col <= rng.max_col:
                anchor = ws.cell(rng.min_row, rng.min_col)
                parsed = parse_portable_label(anchor.value)
                if parsed:
                    return parsed[0], parsed[1], f"{rng}->{anchor.coordinate}"
    return None


def time_for_column(ws, col: int) -> str:
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= 1 <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return clean(ws.cell(rng.min_row, rng.min_col).value)
    return clean(ws.cell(1, col).value)


def extract_portable(ws, exp: int) -> tuple[list[dict], list[str]]:
    errors: list[str] = []
    blocks = []

    for r in range(1, min(6, ws.max_row) + 1):
        for c in range(1, ws.max_column + 1):
            if norm_token(ws.cell(r, c).value) != "ph":
                continue
            dom = find_domain_for_column(ws, r, c)
            if dom is None:
                errors.append(f"EXP{exp}: no domain label for pH block {ws.cell(r,c).coordinate}")
                continue
            raw_r, raw_t, source = dom
            corr_r, corr_t = corrected_domain(exp, raw_r, raw_t)
            blocks.append({
                "header_row": r,
                "start_col": c,
                "raw_r": raw_r,
                "raw_t": raw_t,
                "r": corr_r,
                "t": corr_t,
                "domain": f"R{corr_r}-T{corr_t}",
                "raw_domain": f"R{raw_r}-T{raw_t}",
                "time": time_for_column(ws, c),
                "label_source": source,
            })

    # Parse dates once per physical data row.
    date_records = []
    header_end = max((b["header_row"] for b in blocks), default=3)
    for rr in range(header_end + 1, ws.max_row + 1):
        ts, prov = analytical_portable_date(exp, ws.cell(rr, 1))
        if ts is not None:
            date_records.append({
                "row": rr,
                "date": ts,
                "date_provenance": prov,
                "raw_date": clean(ws.cell(rr, 1).value),
                "date_cell": ws.cell(rr, 1).coordinate,
            })

    repair_excel_date_sequence(date_records)
    date_by_row = {x["row"]: x for x in date_records}

    rows = []
    for b in blocks:
        for rr, dr in date_by_row.items():
            vals = {
                "ph": to_float(ws.cell(rr, b["start_col"]).value),
                "ec": to_float(ws.cell(rr, b["start_col"] + 1).value),
                "tds": to_float(ws.cell(rr, b["start_col"] + 2).value),
                "water_temperature": to_float(ws.cell(rr, b["start_col"] + 3).value),
            }
            # Keep a row if at least one measured value exists.
            if all(pd.isna(v) for v in vals.values()):
                continue
            rows.append({
                "experiment": f"EXP{exp}",
                "experiment_number": exp,
                "domain": b["domain"],
                "replicate": b["r"],
                "treatment": b["t"],
                "raw_portable_domain": b["raw_domain"],
                "raw_portable_replicate": b["raw_r"],
                "raw_portable_treatment": b["raw_t"],
                "time": b["time"],
                "date": dr["date"],
                "raw_date": dr["raw_date"],
                "date_cell": dr["date_cell"],
                "date_provenance": dr["date_provenance"],
                "domain_mapping_provenance": (
                    "P7.6b_EXP1_ROLE_SWAP" if exp == 1 else "RAW_IDENTITY"
                ),
                **vals,
            })

    return rows, errors


def find_total_weight_header_near(ws, domain_row: int) -> tuple[int, int] | None:
    candidates = []
    lo = max(1, domain_row - 4)
    hi = min(ws.max_row, domain_row + 5)
    for r in range(lo, hi + 1):
        for c in range(1, ws.max_column + 1):
            tok = norm_token(ws.cell(r, c).value)
            if tok in {"totalweight", "totalfreshweight", "planttotalweight"}:
                candidates.append((abs(r - domain_row), r, c))
    if not candidates:
        return None
    _, r, c = min(candidates)
    return r, c


def find_plant_number_col(ws, header_row: int) -> int | None:
    for c in range(1, ws.max_column + 1):
        tok = norm_token(ws.cell(header_row, c).value)
        if tok in {"plantno", "plantnumber", "plant", "no"}:
            return c
    return None


def extract_harvest(ws, exp: int) -> tuple[list[dict], list[str]]:
    errors: list[str] = []
    domain_cells = []
    for row in ws.iter_rows():
        for cell in row:
            d = parse_domain_label(cell.value)
            if d:
                domain_cells.append((cell.row, cell.column, d[0], d[1], clean(cell.value)))

    domain_cells.sort()
    rows: list[dict] = []

    for i, (drow, dcol, rep, trt, raw_label) in enumerate(domain_cells):
        nxt = domain_cells[i + 1][0] if i + 1 < len(domain_cells) else ws.max_row + 1
        header = find_total_weight_header_near(ws, drow)
        if header is None:
            errors.append(f"EXP{exp} R{rep}-T{trt}: total-weight header not found near row {drow}")
            continue

        hrow, weight_col = header
        plant_col = find_plant_number_col(ws, hrow)

        # Data begin after both the domain label and the relevant header.
        start = max(drow, hrow) + 1
        stop = min(nxt, ws.max_row + 1)

        candidates = []
        for rr in range(start, stop):
            weight = to_float(ws.cell(rr, weight_col).value)
            if pd.isna(weight):
                continue

            plant_no = to_float(ws.cell(rr, plant_col).value) if plant_col else math.nan
            candidates.append((rr, weight, plant_no))

        # Some layouts put the domain label on the first plant row. If fewer
        # than 12 were found, include the domain row itself when numeric.
        if len(candidates) < EXPECTED_PLANTS_PER_DOMAIN:
            weight = to_float(ws.cell(drow, weight_col).value)
            if not pd.isna(weight):
                plant_no = to_float(ws.cell(drow, plant_col).value) if plant_col else math.nan
                candidates.insert(0, (drow, weight, plant_no))

        # Exclude summary/statistic rows by retaining the first 12 plant-like
        # observations in the domain block.
        candidates = candidates[:EXPECTED_PLANTS_PER_DOMAIN]

        if len(candidates) != EXPECTED_PLANTS_PER_DOMAIN:
            errors.append(
                f"EXP{exp} R{rep}-T{trt}: recovered {len(candidates)} total-weight plants, expected 12"
            )

        for seq, (rr, weight, plant_no) in enumerate(candidates, 1):
            pn = int(plant_no) if not pd.isna(plant_no) and float(plant_no).is_integer() else seq
            rows.append({
                "experiment": f"EXP{exp}",
                "experiment_number": exp,
                "domain": f"R{rep}-T{trt}",
                "replicate": rep,
                "treatment": trt,
                "plant_number": pn,
                "plant_key": f"EXP{exp}_R{rep}_T{trt}_P{pn:02d}",
                "total_weight_g": float(weight),
                "harvest_sheet": ws.title,
                "harvest_source_row": rr,
                "harvest_domain_raw_label": raw_label,
                "total_weight_source_cell": ws.cell(rr, weight_col).coordinate,
            })

    return rows, errors


def slope_per_day(group: pd.DataFrame, variable: str) -> float:
    g = group[["date", variable]].dropna().sort_values("date")
    if len(g) < 2:
        return math.nan
    x = (g["date"] - g["date"].min()).dt.total_seconds().to_numpy() / 86400.0
    y = g[variable].to_numpy(dtype=float)
    if np.ptp(x) == 0:
        return math.nan
    return float(np.polyfit(x, y, 1)[0])


def summarize_environment(long_df: pd.DataFrame) -> pd.DataFrame:
    records = []
    keys = ["experiment", "experiment_number", "domain", "replicate", "treatment"]

    for key, g in long_df.groupby(keys, sort=True):
        rec = dict(zip(keys, key))
        rec["env_start_date"] = g["date"].min().date().isoformat()
        rec["env_end_date"] = g["date"].max().date().isoformat()
        rec["env_unique_dates"] = int(g["date"].dt.normalize().nunique())
        rec["env_measurement_rows"] = int(len(g))
        rec["env_time_groups"] = int(g["time"].nunique())
        rec["env_mapping_provenance"] = " | ".join(sorted(g["domain_mapping_provenance"].unique()))

        for var in VARIABLES:
            s = pd.to_numeric(g[var], errors="coerce").dropna()
            prefix = f"env_{var}"
            rec[f"{prefix}_n"] = int(len(s))
            rec[f"{prefix}_mean"] = float(s.mean()) if len(s) else math.nan
            rec[f"{prefix}_std"] = float(s.std(ddof=1)) if len(s) > 1 else math.nan
            rec[f"{prefix}_min"] = float(s.min()) if len(s) else math.nan
            rec[f"{prefix}_max"] = float(s.max()) if len(s) else math.nan
            rec[f"{prefix}_q25"] = float(s.quantile(0.25)) if len(s) else math.nan
            rec[f"{prefix}_median"] = float(s.median()) if len(s) else math.nan
            rec[f"{prefix}_q75"] = float(s.quantile(0.75)) if len(s) else math.nan
            rec[f"{prefix}_slope"] = slope_per_day(g, var)

        records.append(rec)

    return pd.DataFrame(records).sort_values(keys).reset_index(drop=True)


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    errors: list[str] = []
    warnings: list[str] = []

    if not ARCHIVE.exists():
        print(f"ERROR: archive not found: {ARCHIVE}")
        return 1

    digest = sha256_file(ARCHIVE)
    if digest != EXPECTED_SHA256:
        errors.append("Archive SHA256 does not match the audited Phase-7 archive.")

    all_portable = []
    all_harvest = []
    source_info = []

    with zipfile.ZipFile(ARCHIVE, "r") as zf:
        books = {}
        for member in zf.namelist():
            if member.lower().endswith(".xlsx"):
                exp = experiment_number(member)
                if exp in (1, 2, 3):
                    books[exp] = member

        for exp in (1, 2, 3):
            member = books.get(exp)
            if member is None:
                errors.append(f"EXP{exp}: workbook missing")
                continue

            wb = openpyxl.load_workbook(BytesIO(zf.read(member)), data_only=False)
            portable_name = choose_sheet(wb.sheetnames, "portabl")
            harvest_name = choose_sheet(wb.sheetnames, "harvest")
            if portable_name is None or harvest_name is None:
                errors.append(
                    f"EXP{exp}: missing portable or harvest sheet "
                    f"(portable={portable_name}, harvest={harvest_name})"
                )
                continue

            p_rows, p_errors = extract_portable(wb[portable_name], exp)
            h_rows, h_errors = extract_harvest(wb[harvest_name], exp)
            all_portable.extend(p_rows)
            all_harvest.extend(h_rows)
            errors.extend(p_errors)
            errors.extend(h_errors)
            source_info.append({
                "experiment": f"EXP{exp}",
                "workbook": Path(member).name,
                "portable_sheet": portable_name,
                "harvest_sheet": harvest_name,
                "portable_rows": len(p_rows),
                "harvest_plants": len(h_rows),
            })

    portable_df = pd.DataFrame(all_portable)
    harvest_df = pd.DataFrame(all_harvest)

    if not portable_df.empty:
        portable_df["date"] = pd.to_datetime(portable_df["date"])
        portable_df = portable_df.sort_values(
            ["experiment_number", "domain", "date", "time"]
        ).reset_index(drop=True)
        portable_df.to_csv(PORTABLE_LONG_CSV, index=False)

    if not harvest_df.empty:
        harvest_df = harvest_df.sort_values(
            ["experiment_number", "replicate", "treatment", "plant_number"]
        ).reset_index(drop=True)
        harvest_df.to_csv(HARVEST_CSV, index=False)

    env_df = summarize_environment(portable_df) if not portable_df.empty else pd.DataFrame()
    if not env_df.empty:
        env_df.to_csv(ENV_CSV, index=False)

    join_keys = ["experiment", "experiment_number", "domain", "replicate", "treatment"]
    if not harvest_df.empty and not env_df.empty:
        final_df = harvest_df.merge(
            env_df,
            on=join_keys,
            how="left",
            validate="many_to_one",
            indicator=True,
        )
    else:
        final_df = pd.DataFrame()

    # Integrity checks.
    n_plants = len(final_df)
    n_domains = (
        final_df[["experiment", "domain"]].drop_duplicates().shape[0]
        if not final_df.empty else 0
    )
    duplicate_plant_keys = (
        int(final_df["plant_key"].duplicated().sum()) if not final_df.empty else 0
    )
    unmatched = (
        int((final_df["_merge"] != "both").sum())
        if not final_df.empty and "_merge" in final_df.columns else n_plants
    )
    missing_target = (
        int(final_df["total_weight_g"].isna().sum())
        if not final_df.empty and "total_weight_g" in final_df else n_plants
    )

    per_domain = (
        final_df.groupby(["experiment", "domain"]).size()
        if not final_df.empty else pd.Series(dtype=int)
    )
    bad_domain_sizes = {
        f"{exp}/{dom}": int(n)
        for (exp, dom), n in per_domain.items()
        if int(n) != EXPECTED_PLANTS_PER_DOMAIN
    }

    env_feature_cols = [
        c for c in final_df.columns
        if c.startswith("env_") and any(
            c.endswith("_" + stat) for stat in SUMMARY_STATS
        )
    ] if not final_df.empty else []

    all_missing_env_features = [
        c for c in env_feature_cols if final_df[c].isna().all()
    ] if env_feature_cols else []

    # Verify documented EXP2 correction appears in reconstructed long data.
    exp2_corrections = 0
    if not portable_df.empty:
        exp2_corrections = int(
            portable_df["date_provenance"]
            .astype(str)
            .str.contains("DOCUMENTED_ANALYTICAL_CORRECTION")
            .sum()
        )
        # The same physical corrected date is repeated across 12 portable
        # domain/time blocks, so presence rather than count is the safeguard.
        if exp2_corrections == 0:
            errors.append("Documented EXP2 A31/A32 analytical correction was not represented.")

    # Check environmental domain coverage.
    env_domains = (
        env_df[["experiment", "domain"]].drop_duplicates().shape[0]
        if not env_df.empty else 0
    )

    # Date monotonicity is checked at the experiment/date-cell level, not after
    # duplication across domains.
    temporal_issues = []
    if not portable_df.empty:
        for exp, g in portable_df.groupby("experiment"):
            d = (
                g[["date_cell", "date"]]
                .drop_duplicates("date_cell")
                .sort_values("date_cell", key=lambda s: s.str.extract(r"(\d+)")[0].astype(int))
            )
            diffs = d["date"].diff().dropna().dt.days
            backwards = int((diffs < 0).sum())
            if backwards:
                temporal_issues.append(f"{exp}: {backwards} backward date step(s)")

    if temporal_issues:
        errors.extend(temporal_issues)

    if n_plants != EXPECTED_PLANTS:
        errors.append(f"Final plant count {n_plants}, expected {EXPECTED_PLANTS}.")
    if n_domains != EXPECTED_DOMAINS:
        errors.append(f"Final domain count {n_domains}, expected {EXPECTED_DOMAINS}.")
    if env_domains != EXPECTED_DOMAINS:
        errors.append(f"Environmental domain count {env_domains}, expected {EXPECTED_DOMAINS}.")
    if duplicate_plant_keys:
        errors.append(f"Duplicate plant keys: {duplicate_plant_keys}.")
    if unmatched:
        errors.append(f"Plants without environmental domain match: {unmatched}.")
    if missing_target:
        errors.append(f"Plants with missing primary target: {missing_target}.")
    if bad_domain_sizes:
        errors.append(f"Domains not containing 12 plants: {bad_domain_sizes}.")
    if all_missing_env_features:
        errors.append(f"All-missing environmental features: {all_missing_env_features}.")

    # Remove merge helper only after diagnostics.
    if not final_df.empty and "_merge" in final_df.columns:
        final_df = final_df.drop(columns=["_merge"])

    if not final_df.empty:
        final_df.to_csv(DATASET_CSV, index=False)

    verdict = "PASS" if not errors else "REVIEW REQUIRED"

    report = [
        "PHASE 7.8 — PLANT-LEVEL ANALYTICAL DATASET CONSTRUCTION",
        "=" * 118,
        f"Archive: {ARCHIVE}",
        f"Archive SHA256: {digest}",
        f"Expected SHA256 match: {'YES' if digest == EXPECTED_SHA256 else 'NO'}",
        "",
        "FROZEN P7.7 DESIGN IMPLEMENTED",
        "-" * 118,
        "Experimental observation: individual harvested lettuce plant",
        "Primary target: final total fresh plant weight (g)",
        "Environmental unit: experiment × treatment × replicate domain",
        "Primary environmental histories: pH, EC, TDS, water temperature",
        "History representation: full observed portable experiment history; no invented harvest-relative window",
        "EXP1 domain correction: P7.6b role swap, analytical mapping only",
        "EXP2 date correction: A31/A32 interpreted analytically as 13/05/2024 and 14/05/2024",
        "Synthetic oracle augmentation: NOT APPLIED",
        "Model training: NOT PERFORMED",
        "",
        "SOURCE EXTRACTION",
        "-" * 118,
    ]

    for x in source_info:
        report.extend([
            f"{x['experiment']}:",
            f"  workbook={x['workbook']!r}",
            f"  portable_sheet={x['portable_sheet']!r}",
            f"  harvest_sheet={x['harvest_sheet']!r}",
            f"  portable measurement rows retained={x['portable_rows']}",
            f"  harvested plants recovered={x['harvest_plants']}",
        ])

    report.extend([
        "",
        "FINAL DATASET INTEGRITY",
        "-" * 118,
        f"Plant rows: {n_plants} / expected {EXPECTED_PLANTS}",
        f"Unique experiment-domain cells: {n_domains} / expected {EXPECTED_DOMAINS}",
        f"Environmental domain summaries: {env_domains} / expected {EXPECTED_DOMAINS}",
        f"Duplicate plant keys: {duplicate_plant_keys}",
        f"Plants without environmental-domain match: {unmatched}",
        f"Missing primary targets: {missing_target}",
        f"Domains with plant count != 12: {len(bad_domain_sizes)}",
        f"Environmental summary predictor columns: {len(env_feature_cols)}",
        f"All-missing environmental predictor columns: {len(all_missing_env_features)}",
        f"EXP2 corrected-date measurement rows represented: {exp2_corrections}",
        f"Temporal backward-step issues after analytical reconstruction: {len(temporal_issues)}",
        "",
        "LEAKAGE / UNIT-OF-ANALYSIS SAFEGUARDS",
        "-" * 118,
        "Portable measurement rows are summarized to environmental histories.",
        "Portable/hourly rows are NOT exported as independent plant training observations.",
        "Each environmental history joins many-to-one to plants in exactly one experiment-domain cell.",
        "Future OOD evaluation must split by held-out domain, not random plant rows across all domains.",
        "",
        "VALIDATION ERRORS",
        "-" * 118,
        f"Error count: {len(errors)}",
    ])
    report.extend([f"ERROR: {e}" for e in errors] or ["None"])

    report.extend([
        "",
        "FINAL VERDICT",
        "=" * 118,
        verdict,
    ])

    if verdict == "PASS":
        report.extend([
            "The 216-plant analytical table was constructed with complete target/domain linkage.",
            "Environmental histories are represented at domain level and attached to plants without treating sensor rows as independent observations.",
            "This PASS authorizes modeling-protocol implementation next; it does not itself report ML performance.",
        ])

    report.extend([
        "",
        "RESEARCH LOCK STATUS",
        "-" * 118,
        "P7.7 analytical design implemented: YES",
        f"P7.8 plant-level dataset validated: {'YES' if verdict == 'PASS' else 'NO'}",
        "Raw archive modified: NO",
        "Primary target used: total_weight_g",
        "Synthetic oracle augmentation applied: NO",
        "Model training performed: NO",
    ])

    REPORT.write_text("\n".join(report) + "\n", encoding="utf-8")

    manifest = {
        "phase": "P7.8",
        "archive_sha256": digest,
        "expected_sha256_match": digest == EXPECTED_SHA256,
        "expected_plants": EXPECTED_PLANTS,
        "actual_plants": n_plants,
        "expected_domains": EXPECTED_DOMAINS,
        "actual_domains": n_domains,
        "environmental_domains": env_domains,
        "duplicate_plant_keys": duplicate_plant_keys,
        "unmatched_plants": unmatched,
        "missing_primary_targets": missing_target,
        "bad_domain_sizes": bad_domain_sizes,
        "environmental_feature_count": len(env_feature_cols),
        "all_missing_environmental_features": all_missing_env_features,
        "exp2_corrected_date_measurement_rows": exp2_corrections,
        "temporal_issues": temporal_issues,
        "errors": errors,
        "warnings": warnings,
        "verdict": verdict,
        "raw_archive_modified": False,
        "model_training_performed": False,
        "outputs": {
            "dataset": str(DATASET_CSV),
            "environment_summaries": str(ENV_CSV),
            "harvest_outcomes": str(HARVEST_CSV),
            "portable_long": str(PORTABLE_LONG_CSV),
            "report": str(REPORT),
        },
    }
    MANIFEST.write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )

    print("\n".join(report))
    print("")
    print(f"Dataset: {DATASET_CSV}")
    print(f"Report: {REPORT}")
    print("RAW DATA WERE NOT MODIFIED.")
    print("NO MODEL TRAINING WAS PERFORMED.")
    return 0 if verdict == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
