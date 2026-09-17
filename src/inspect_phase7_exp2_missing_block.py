#!/usr/bin/env python3
"""
Phase 7.6d — EXP2 Missing Portable-Block Forensic Inspection

Purpose
-------
Explain why P7.6c recovered 11 rather than 12 EXP2 portable measurement blocks.

This diagnostic compares:
1. all raw EXP2 portable merged domain labels,
2. all pH header positions,
3. expected 10:00 / 14:00 domain slots,
4. the P7.6c upward-label detector,
5. a geometry-based detector using merged ranges.

It does NOT modify raw files, correct mappings, train models, or freeze Phase 7.
"""

from __future__ import annotations

import hashlib
import json
import re
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
ARCHIVE = ROOT / "data" / "phase7" / "raw" / "all_months_sensory_data.zip"
OUT = ROOT / "results" / "experiment_7" / "exp2_missing_block"

REPORT = OUT / "phase7_exp2_missing_block_report.txt"
SLOTS_CSV = OUT / "phase7_exp2_expected_slots.csv"
HEADERS_CSV = OUT / "phase7_exp2_header_evidence.csv"
MERGES_CSV = OUT / "phase7_exp2_merged_domain_labels.csv"
MANIFEST = OUT / "phase7_exp2_missing_block_manifest.json"

EXPECTED_SHA256 = "daa6d1cc259234d4b1cb8ba9901b2f7e35975a444154d5b8dd0a58a394fdbf14"


def clean(v: Any) -> str:
    if v is None:
        return ""
    try:
        if pd.isna(v):
            return ""
    except Exception:
        pass
    return re.sub(r"\s+", " ", str(v)).strip()


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def exp_number(name: str) -> int | None:
    m = re.search(r"EXP\.?\s*(\d+)", Path(name).name, flags=re.I)
    return int(m.group(1)) if m else None


def choose_sheet(names: list[str], hint: str) -> str | None:
    for name in names:
        if hint.lower() in name.lower():
            return name
    return None


def parse_label(v: Any) -> tuple[int, int] | None:
    s = clean(v).upper()
    r = re.search(r"\bREPLICATE\s*[-_:]?\s*(\d+)\b", s)
    t = re.search(r"\bT\s*[-_:]?\s*(\d+)\b", s)
    if r and t:
        return int(r.group(1)), int(t.group(1))
    return None


def merged_domain_labels(ws) -> list[dict]:
    out = []
    for rng in ws.merged_cells.ranges:
        anchor = ws.cell(rng.min_row, rng.min_col)
        parsed = parse_label(anchor.value)
        if parsed:
            r, t = parsed
            out.append({
                "range": str(rng),
                "anchor": anchor.coordinate,
                "value": clean(anchor.value),
                "replicate": r,
                "treatment": t,
                "domain": f"R{r}-T{t}",
                "min_row": rng.min_row,
                "max_row": rng.max_row,
                "min_col": rng.min_col,
                "max_col": rng.max_col,
            })
    return sorted(out, key=lambda x: (x["min_row"], x["min_col"]))


def pH_headers(ws) -> list[dict]:
    out = []
    for r in range(1, min(6, ws.max_row) + 1):
        for c in range(1, ws.max_column + 1):
            if clean(ws.cell(r, c).value).lower() == "ph":
                out.append({
                    "row": r,
                    "col": c,
                    "cell": ws.cell(r, c).coordinate,
                    "start_col": get_column_letter(c),
                    "end_col": get_column_letter(min(c + 3, ws.max_column)),
                    "variables": " | ".join(
                        clean(ws.cell(r, cc).value)
                        for cc in range(c, min(c + 3, ws.max_column) + 1)
                    ),
                })
    return out


def time_for_column(ws, col: int) -> str:
    direct = clean(ws.cell(1, col).value)
    if direct:
        return direct
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= 1 <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return clean(ws.cell(rng.min_row, rng.min_col).value)
    return ""


def p76c_detector(ws, ph_row: int, col: int) -> tuple[str, str]:
    """Reproduce the label-detection logic used in P7.6c."""
    for hr in range(ph_row - 1, 0, -1):
        v = ws.cell(hr, col).value
        parsed = parse_label(v)
        if parsed:
            r, t = parsed
            return f"R{r}-T{t}", f"direct:{ws.cell(hr,col).coordinate}"

        for rng in ws.merged_cells.ranges:
            if (
                rng.min_row <= hr <= rng.max_row
                and rng.min_col <= col <= rng.max_col
            ):
                anchor = ws.cell(rng.min_row, rng.min_col)
                parsed = parse_label(anchor.value)
                if parsed:
                    r, t = parsed
                    return f"R{r}-T{t}", f"merged:{rng}->{anchor.coordinate}"
    return "", ""


def geometry_detector(
    merges: list[dict],
    ph_row: int,
    col: int,
) -> tuple[str, str]:
    """
    Associate a pH block with a domain label whose merged range geometrically
    covers its starting column. This avoids relying on a non-empty anchor being
    located exactly above the pH cell.
    """
    candidates = [
        m for m in merges
        if m["min_col"] <= col <= m["max_col"] and m["max_row"] < ph_row
    ]
    if not candidates:
        return "", ""
    candidates.sort(key=lambda x: (ph_row - x["max_row"], x["min_col"]))
    m = candidates[0]
    return m["domain"], f"{m['range']}->{m['anchor']}"


def top_grid(ws, rows: int = 5) -> list[str]:
    lines = []
    for r in range(1, min(rows, ws.max_row) + 1):
        vals = []
        for c in range(1, ws.max_column + 1):
            v = clean(ws.cell(r, c).value)
            if v:
                vals.append(f"{ws.cell(r,c).coordinate}={v!r}")
        lines.append(f"Row {r}: " + (" | ".join(vals) if vals else "<empty>"))
    return lines


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    errors: list[str] = []

    report = [
        "PHASE 7.6d — EXP2 MISSING PORTABLE-BLOCK FORENSIC INSPECTION",
        "=" * 118,
        f"Archive: {ARCHIVE}",
        "Purpose: explain why P7.6c recovered 11 rather than 12 EXP2 portable measurement blocks.",
        "No raw mutation, no mapping correction, no model training, no Phase-7 design freeze.",
        "",
    ]

    if not ARCHIVE.exists():
        print(f"ERROR: archive not found: {ARCHIVE}")
        return 1

    digest = sha256_file(ARCHIVE)
    report.extend([
        f"Archive SHA256: {digest}",
        f"Expected SHA256 match: {'YES' if digest == EXPECTED_SHA256 else 'NO'}",
        "",
    ])

    with zipfile.ZipFile(ARCHIVE, "r") as zf:
        members = [
            n for n in zf.namelist()
            if n.lower().endswith(".xlsx") and exp_number(n) == 2
        ]
        if len(members) != 1:
            errors.append(f"Expected exactly one EXP2 workbook; found {len(members)}: {members}")
            REPORT.write_text("\n".join(report + errors) + "\n", encoding="utf-8")
            return 1
        member = members[0]
        raw = zf.read(member)

    wb = openpyxl.load_workbook(BytesIO(raw), data_only=False)
    portable_name = choose_sheet(wb.sheetnames, "portabl")
    if portable_name is None:
        print("ERROR: EXP2 portable sheet not found")
        return 1
    ws = wb[portable_name]

    merges = merged_domain_labels(ws)
    phs = pH_headers(ws)

    report.extend([
        "SOURCE DISCOVERY",
        "-" * 118,
        f"Workbook: {Path(member).name!r}",
        f"Portable sheet: {portable_name!r}",
        f"Worksheet shape: {ws.max_row} rows × {ws.max_column} columns",
        f"Merged domain-label ranges detected: {len(merges)}",
        f"pH headers detected: {len(phs)}",
        "",
        "RAW TOP-HEADER GRID",
        "-" * 118,
    ])
    report.extend(top_grid(ws, 5))
    report.append("")

    report.extend([
        "MERGED DOMAIN LABELS",
        "-" * 118,
    ])
    for m in merges:
        report.append(
            f"{m['range']}: {m['value']!r} -> {m['domain']} "
            f"(cols {m['min_col']}..{m['max_col']})"
        )
    report.append("")

    slot_rows = []
    header_rows = []

    for p in phs:
        time_label = time_for_column(ws, p["col"])
        old_domain, old_src = p76c_detector(ws, p["row"], p["col"])
        geo_domain, geo_src = geometry_detector(merges, p["row"], p["col"])

        row = {
            "time_label": time_label,
            "ph_cell": p["cell"],
            "block_columns": f"{p['start_col']}:{p['end_col']}",
            "variables": p["variables"],
            "p76c_domain": old_domain,
            "p76c_source": old_src,
            "geometry_domain": geo_domain,
            "geometry_source": geo_src,
            "detectors_agree": bool(old_domain) and old_domain == geo_domain,
        }
        slot_rows.append(row)
        header_rows.append(row.copy())

    pd.DataFrame(slot_rows).to_csv(SLOTS_CSV, index=False)
    pd.DataFrame(header_rows).to_csv(HEADERS_CSV, index=False)
    pd.DataFrame(merges).to_csv(MERGES_CSV, index=False)

    report.extend([
        "BLOCK-BY-BLOCK DETECTOR COMPARISON",
        "-" * 118,
    ])
    for i, x in enumerate(slot_rows, 1):
        report.append(
            f"Block {i}: time={x['time_label']!r}, ph={x['ph_cell']}, "
            f"cols={x['block_columns']}, vars={x['variables']}"
        )
        report.append(
            f"  P7.6c detector: {x['p76c_domain'] or 'UNRESOLVED'} "
            f"via {x['p76c_source'] or '-'}"
        )
        report.append(
            f"  Geometry detector: {x['geometry_domain'] or 'UNRESOLVED'} "
            f"via {x['geometry_source'] or '-'}"
        )
    report.append("")

    # Expected domain x time combinations inferred from source labels/time groups.
    raw_domains = sorted({m["domain"] for m in merges})
    times = sorted({x["time_label"] for x in slot_rows if x["time_label"]})

    expected = {(t, d) for t in times for d in raw_domains}
    old_found = {
        (x["time_label"], x["p76c_domain"])
        for x in slot_rows
        if x["time_label"] and x["p76c_domain"]
    }
    geo_found = {
        (x["time_label"], x["geometry_domain"])
        for x in slot_rows
        if x["time_label"] and x["geometry_domain"]
    }

    missing_old = sorted(expected - old_found)
    missing_geo = sorted(expected - geo_found)
    extra_old = sorted(old_found - expected)
    extra_geo = sorted(geo_found - expected)

    report.extend([
        "EXPECTED-SLOT ACCOUNTING",
        "-" * 118,
        f"Distinct source domains: {raw_domains}",
        f"Distinct time groups: {times}",
        f"Expected domain×time slots: {len(expected)}",
        f"P7.6c resolved unique slots: {len(old_found)}",
        f"Geometry resolved unique slots: {len(geo_found)}",
        f"P7.6c missing slots: {missing_old}",
        f"Geometry missing slots: {missing_geo}",
        f"P7.6c unexpected slots: {extra_old}",
        f"Geometry unexpected slots: {extra_geo}",
        "",
    ])

    unresolved_old = [x for x in slot_rows if not x["p76c_domain"]]
    unresolved_geo = [x for x in slot_rows if not x["geometry_domain"]]

    if len(phs) == 12 and len(expected) == 12 and not missing_geo:
        if missing_old or unresolved_old:
            status = "PARSER_DETECTION_ISSUE_CONFIRMED"
            interpretation = (
                "The raw worksheet contains all 12 pH measurement blocks and the geometry-based "
                "header linkage resolves all 12 expected domain×time slots. The P7.6c 11-block "
                "count is therefore a parser-detection issue, not a missing source-data block."
            )
        else:
            status = "ALL_12_BLOCKS_RESOLVED"
            interpretation = (
                "Both detectors resolve all 12 expected slots. The prior 11-block result requires "
                "comparison with the exact P7.6c implementation/output."
            )
    elif len(phs) < 12:
        status = "SOURCE_OR_HEADER_BLOCK_MISSING"
        interpretation = (
            f"Only {len(phs)} pH headers were found, so the source/header layout itself may lack "
            "an expected measurement block."
        )
    else:
        status = "UNRESOLVED"
        interpretation = (
            "The geometry-based audit did not recover all expected domain×time slots; further "
            "source-level investigation is required."
        )

    report.extend([
        "FORENSIC INTERPRETATION",
        "-" * 118,
        interpretation,
        f"Forensic status: {status}",
        "",
        "VALIDATION",
        "-" * 118,
        f"Processing errors: {len(errors)}",
        f"Archive hash verified: {'YES' if digest == EXPECTED_SHA256 else 'NO'}",
        f"pH blocks found: {len(phs)}",
        f"Expected slots: {len(expected)}",
        f"P7.6c unresolved blocks in this reproduction: {len(unresolved_old)}",
        f"Geometry unresolved blocks: {len(unresolved_geo)}",
        f"Geometry missing expected slots: {len(missing_geo)}",
    ])
    for e in errors:
        report.append(f"ERROR: {e}")

    report.extend([
        "",
        "RESEARCH LOCK STATUS",
        "-" * 118,
        "EXP2 source block corrected: NO",
        "P7.6c parser changed: NO",
        "Cross-table identity-level linkage frozen: NO",
        "Primary target frozen: NO",
        "Feature set frozen: NO",
        "Domain definition frozen: NO",
        "Shift/severity frozen: NO",
        "Model training performed: NO",
    ])

    REPORT.write_text("\n".join(report) + "\n", encoding="utf-8")

    manifest = {
        "phase": "P7.6d",
        "archive_sha256": digest,
        "expected_sha256_match": digest == EXPECTED_SHA256,
        "workbook": Path(member).name,
        "portable_sheet": portable_name,
        "ph_blocks_found": len(phs),
        "expected_slots": len(expected),
        "p76c_slots_found": len(old_found),
        "geometry_slots_found": len(geo_found),
        "p76c_missing_slots": missing_old,
        "geometry_missing_slots": missing_geo,
        "forensic_status": status,
        "raw_modified": False,
        "parser_changed": False,
        "model_training_performed": False,
        "outputs": {
            "report": str(REPORT),
            "slots": str(SLOTS_CSV),
            "headers": str(HEADERS_CSV),
            "merged_labels": str(MERGES_CSV),
        },
    }
    MANIFEST.write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    print("\n".join(report))
    print("")
    print(f"Report: {REPORT}")
    print("RAW DATA WERE NOT MODIFIED.")
    print("NO MODEL TRAINING WAS PERFORMED.")
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
