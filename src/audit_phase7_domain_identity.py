#!/usr/bin/env python3
"""
Phase 7.6 — Cross-table domain identity audit

Purpose
-------
Audit whether treatment/replicate domain identities reconstructed from the
harvest sheets agree with identities represented in the portable environmental
sheets, with special attention to EXP1.

This script is DIAGNOSTIC ONLY:
- does not modify the raw archive
- does not train models
- does not impute data
- does not freeze target/features/domains/shifts/models
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd


ROOT = Path(__file__).resolve().parents[1]
ARCHIVE = ROOT / "data" / "phase7" / "raw" / "all_months_sensory_data.zip"
OUTDIR = ROOT / "results" / "experiment_7" / "domain_identity_audit"

REPORT = OUTDIR / "phase7_domain_identity_report.txt"
COMPARISON_CSV = OUTDIR / "phase7_domain_identity_comparison.csv"
LABELS_CSV = OUTDIR / "phase7_raw_domain_labels.csv"
HEADERS_CSV = OUTDIR / "phase7_portable_header_evidence.csv"
MANIFEST = OUTDIR / "phase7_domain_identity_manifest.json"

EXPECTED_SHA256 = "daa6d1cc259234d4b1cb8ba9901b2f7e35975a444154d5b8dd0a58a394fdbf14"


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def norm_text(v: Any) -> str:
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v).strip())


def compact(v: Any) -> str:
    return re.sub(r"[^A-Z0-9]+", "", norm_text(v).upper())


def domain_from_text(v: Any) -> str | None:
    """Conservatively recognize explicit R#T# / R#-T# domain labels."""
    s = compact(v)
    m = re.fullmatch(r"R(\d+)T(\d+)", s)
    if not m:
        return None
    return f"R{int(m.group(1))}-T{int(m.group(2))}"


def treatment_from_text(v: Any) -> int | None:
    s = norm_text(v).upper()
    patterns = [
        r"\bT(?:REATMENT)?\s*[-_:]?\s*(\d+)\b",
        r"\bTREATMENT\s*(\d+)\b",
    ]
    for p in patterns:
        m = re.search(p, s)
        if m:
            return int(m.group(1))
    return None


def replicate_from_text(v: Any) -> int | None:
    s = norm_text(v).upper()
    patterns = [
        r"\bR(?:EPLICATE)?\s*[-_:]?\s*(\d+)\b",
        r"\bREPLICATE\s*(\d+)\b",
    ]
    for p in patterns:
        m = re.search(p, s)
        if m:
            return int(m.group(1))
    return None


def experiment_number(name: str) -> int | None:
    m = re.search(r"EXP\.?\s*(\d+)", name, flags=re.I)
    return int(m.group(1)) if m else None


def find_workbooks(extract_dir: Path) -> dict[int, Path]:
    result: dict[int, Path] = {}
    for p in sorted(extract_dir.rglob("*.xlsx")):
        n = experiment_number(p.name)
        if n in (1, 2, 3):
            result[n] = p
    return result


def choose_sheet(wb: openpyxl.Workbook, keyword: str) -> str | None:
    k = keyword.lower()
    for name in wb.sheetnames:
        if k in name.lower():
            return name
    return None


def harvest_domain_evidence(ws, exp: int) -> tuple[list[dict], set[str]]:
    rows: list[dict] = []
    domains: set[str] = set()

    for row in ws.iter_rows():
        for cell in row:
            d = domain_from_text(cell.value)
            if d:
                domains.add(d)
                rows.append({
                    "experiment": f"EXP{exp}",
                    "source": "harvest",
                    "sheet": ws.title,
                    "cell": cell.coordinate,
                    "raw_value": norm_text(cell.value),
                    "parsed_domain": d,
                    "evidence_type": "explicit_RxTx_label",
                })
    return rows, domains


def merged_anchor_value(ws, row: int, col: int) -> Any:
    """Return value at cell or merged-range anchor if the cell belongs to a merge."""
    c = ws.cell(row=row, column=col)
    if c.value is not None:
        return c.value
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return ws.cell(rng.min_row, rng.min_col).value
    return None


def portable_header_evidence(ws, exp: int) -> tuple[list[dict], list[dict], set[str]]:
    """
    Inspect top header rows without assuming EXP1 has the same orientation as EXP2/3.

    We record raw/merged header evidence and derive a domain only when both an
    explicit treatment number and replicate number can be associated with a
    4-column pH/EC/TDS/water-temperature measurement block.
    """
    raw_rows: list[dict] = []
    header_rows: list[dict] = []
    domains: set[str] = set()

    max_header_row = min(5, ws.max_row)

    # Catalog all non-empty top-header cells and merged ranges.
    for r in range(1, max_header_row + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if v is not None and norm_text(v):
                raw_rows.append({
                    "experiment": f"EXP{exp}",
                    "source": "portable",
                    "sheet": ws.title,
                    "cell": ws.cell(r, c).coordinate,
                    "raw_value": norm_text(v),
                    "parsed_domain": domain_from_text(v) or "",
                    "evidence_type": "top_header_cell",
                })

    for rng in ws.merged_cells.ranges:
        if rng.min_row <= max_header_row:
            v = ws.cell(rng.min_row, rng.min_col).value
            header_rows.append({
                "experiment": f"EXP{exp}",
                "sheet": ws.title,
                "range_or_cell": str(rng),
                "row": rng.min_row,
                "start_col": rng.min_col,
                "end_col": rng.max_col,
                "value": norm_text(v),
                "treatment_number": treatment_from_text(v),
                "replicate_number": replicate_from_text(v),
                "kind": "merged_range",
            })

    # Add ordinary header cells as evidence too.
    for r in range(1, max_header_row + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if v is not None and norm_text(v):
                header_rows.append({
                    "experiment": f"EXP{exp}",
                    "sheet": ws.title,
                    "range_or_cell": ws.cell(r, c).coordinate,
                    "row": r,
                    "start_col": c,
                    "end_col": c,
                    "value": norm_text(v),
                    "treatment_number": treatment_from_text(v),
                    "replicate_number": replicate_from_text(v),
                    "kind": "cell",
                })

    # Locate pH starts in first five rows. Each domain block in this workbook
    # family is expected to expose pH/EC/TDS/water-temp across four columns.
    ph_starts: list[tuple[int, int]] = []
    for r in range(1, max_header_row + 1):
        for c in range(1, ws.max_column + 1):
            s = norm_text(ws.cell(r, c).value).lower()
            if s == "ph" or s.startswith("ph "):
                ph_starts.append((r, c))

    for ph_row, start_col in ph_starts:
        # Search header context covering the 4-column block.
        treatments: list[int] = []
        replicates: list[int] = []
        evidence: list[str] = []

        for rr in range(1, max_header_row + 1):
            for cc in range(start_col, min(start_col + 3, ws.max_column) + 1):
                v = merged_anchor_value(ws, rr, cc)
                if v is None:
                    continue
                t = treatment_from_text(v)
                rep = replicate_from_text(v)
                if t is not None:
                    treatments.append(t)
                    evidence.append(f"{ws.cell(rr, cc).coordinate}:{norm_text(v)}")
                if rep is not None:
                    replicates.append(rep)
                    evidence.append(f"{ws.cell(rr, cc).coordinate}:{norm_text(v)}")

        # Also inspect merged ranges that geometrically cover this block.
        for rng in ws.merged_cells.ranges:
            overlap = not (rng.max_col < start_col or rng.min_col > start_col + 3)
            if overlap and rng.min_row <= max_header_row:
                v = ws.cell(rng.min_row, rng.min_col).value
                t = treatment_from_text(v)
                rep = replicate_from_text(v)
                if t is not None:
                    treatments.append(t)
                    evidence.append(f"{rng}:{norm_text(v)}")
                if rep is not None:
                    replicates.append(rep)
                    evidence.append(f"{rng}:{norm_text(v)}")

        tset = sorted(set(treatments))
        rset = sorted(set(replicates))
        parsed = ""
        status = "AMBIGUOUS"

        if len(tset) == 1 and len(rset) == 1:
            parsed = f"R{rset[0]}-T{tset[0]}"
            domains.add(parsed)
            status = "RESOLVED"

        header_rows.append({
            "experiment": f"EXP{exp}",
            "sheet": ws.title,
            "range_or_cell": f"measurement_block_{start_col}:{start_col+3}",
            "row": ph_row,
            "start_col": start_col,
            "end_col": min(start_col + 3, ws.max_column),
            "value": " | ".join(sorted(set(evidence))),
            "treatment_number": tset[0] if len(tset) == 1 else None,
            "replicate_number": rset[0] if len(rset) == 1 else None,
            "kind": f"measurement_block_{status}_{parsed}",
        })

    return raw_rows, header_rows, domains


def shape_signature(domains: set[str]) -> tuple[int, int, list[int], list[int]]:
    reps, treats = set(), set()
    for d in domains:
        m = re.fullmatch(r"R(\d+)-T(\d+)", d)
        if m:
            reps.add(int(m.group(1)))
            treats.add(int(m.group(2)))
    return len(reps), len(treats), sorted(reps), sorted(treats)


def main() -> int:
    OUTDIR.mkdir(parents=True, exist_ok=True)

    lines: list[str] = []
    errors: list[str] = []

    lines.append("PHASE 7.6 — CROSS-TABLE DOMAIN IDENTITY AUDIT")
    lines.append("=" * 118)
    lines.append(f"Archive: {ARCHIVE}")
    lines.append("Purpose: compare harvest-domain identities with portable-environment domain identities.")
    lines.append("No raw mutation, no model training, no Phase-7 design freeze.")
    lines.append("")

    if not ARCHIVE.exists():
        print(f"ERROR: archive not found: {ARCHIVE}")
        return 1

    digest = sha256_file(ARCHIVE)
    lines.append(f"Archive SHA256: {digest}")
    lines.append(f"Expected SHA256 match: {'YES' if digest == EXPECTED_SHA256 else 'NO'}")
    lines.append("")

    extract_dir = OUTDIR / "_extracted"
    if extract_dir.exists():
        for p in sorted(extract_dir.rglob("*"), reverse=True):
            if p.is_file():
                p.unlink()
            elif p.is_dir():
                try:
                    p.rmdir()
                except OSError:
                    pass
    extract_dir.mkdir(parents=True, exist_ok=True)

    with zipfile.ZipFile(ARCHIVE) as zf:
        zf.extractall(extract_dir)

    workbooks = find_workbooks(extract_dir)

    all_labels: list[dict] = []
    all_headers: list[dict] = []
    comparisons: list[dict] = []

    lines.append("WORKBOOK / SHEET DISCOVERY")
    lines.append("-" * 118)

    for exp in (1, 2, 3):
        path = workbooks.get(exp)
        if path is None:
            errors.append(f"EXP{exp}: workbook not found")
            continue

        try:
            wb = openpyxl.load_workbook(path, data_only=False)
        except Exception as e:
            errors.append(f"EXP{exp}: workbook read failed: {type(e).__name__}: {e}")
            continue

        harvest_name = choose_sheet(wb, "harvest")
        portable_name = choose_sheet(wb, "portabl")

        lines.append(
            f"EXP{exp}: workbook={path.name!r}; harvest={harvest_name!r}; portable={portable_name!r}"
        )

        if harvest_name is None or portable_name is None:
            errors.append(f"EXP{exp}: required harvest/portable sheet missing")
            continue

        harvest_rows, harvest_domains = harvest_domain_evidence(wb[harvest_name], exp)
        portable_rows, header_rows, portable_domains = portable_header_evidence(
            wb[portable_name], exp
        )

        all_labels.extend(harvest_rows)
        all_labels.extend(portable_rows)
        all_headers.extend(header_rows)

        h_reps_n, h_treats_n, h_reps, h_treats = shape_signature(harvest_domains)
        p_reps_n, p_treats_n, p_reps, p_treats = shape_signature(portable_domains)

        union = sorted(harvest_domains | portable_domains)
        for d in union:
            comparisons.append({
                "experiment": f"EXP{exp}",
                "domain": d,
                "in_harvest": d in harvest_domains,
                "in_portable": d in portable_domains,
                "status": "MATCH" if d in harvest_domains and d in portable_domains else "MISMATCH",
            })

        lines.append("")
        lines.append(f"EXP{exp} DOMAIN IDENTITY")
        lines.append("-" * 118)
        lines.append(f"Harvest explicit domains ({len(harvest_domains)}): {sorted(harvest_domains)}")
        lines.append(f"Portable resolved domains ({len(portable_domains)}): {sorted(portable_domains)}")
        lines.append(
            f"Harvest shape: {h_reps_n} replicate(s) {h_reps} × "
            f"{h_treats_n} treatment(s) {h_treats}"
        )
        lines.append(
            f"Portable shape: {p_reps_n} replicate(s) {p_reps} × "
            f"{p_treats_n} treatment(s) {p_treats}"
        )
        lines.append(f"Harvest-only: {sorted(harvest_domains - portable_domains)}")
        lines.append(f"Portable-only: {sorted(portable_domains - harvest_domains)}")
        lines.append(
            "Exact domain-set agreement: "
            + ("YES" if harvest_domains and harvest_domains == portable_domains else "NO")
        )

    pd.DataFrame(all_labels).to_csv(LABELS_CSV, index=False)
    pd.DataFrame(all_headers).to_csv(HEADERS_CSV, index=False)
    pd.DataFrame(comparisons).to_csv(COMPARISON_CSV, index=False)

    lines.append("")
    lines.append("CROSS-TABLE SUMMARY")
    lines.append("-" * 118)

    comp_df = pd.DataFrame(comparisons)
    if not comp_df.empty:
        for exp in ("EXP1", "EXP2", "EXP3"):
            sub = comp_df[comp_df["experiment"] == exp]
            counts = Counter(sub["status"])
            lines.append(
                f"{exp}: compared domain keys={len(sub)}, "
                f"MATCH={counts.get('MATCH', 0)}, MISMATCH={counts.get('MISMATCH', 0)}"
            )

    exp1_mismatch = False
    if not comp_df.empty:
        e1 = comp_df[comp_df["experiment"] == "EXP1"]
        exp1_mismatch = bool((e1["status"] == "MISMATCH").any())

    lines.append("")
    lines.append("INTERPRETATION")
    lines.append("-" * 118)
    if exp1_mismatch:
        lines.append(
            "EXP1 domain identities do NOT align exactly between harvest and portable environmental evidence."
        )
        lines.append(
            "Therefore the earlier count-based 18/18 linkage must NOT be treated as proof of identity-level linkage."
        )
        lines.append(
            "Inspect phase7_portable_header_evidence.csv before changing any EXP1 parser or domain mapping."
        )
    else:
        lines.append(
            "No EXP1 identity mismatch was detected by this conservative raw-header audit."
        )
        lines.append(
            "This supports identity-level linkage only to the extent that the raw headers were explicitly resolvable."
        )

    lines.append("")
    lines.append("VALIDATION")
    lines.append("-" * 118)
    lines.append(f"Processing errors: {len(errors)}")
    for e in errors:
        lines.append(f"ERROR: {e}")
    lines.append(f"Workbooks found: {len(workbooks)}")
    lines.append(f"Raw label evidence rows: {len(all_labels)}")
    lines.append(f"Portable header evidence rows: {len(all_headers)}")

    structural_fail = bool(errors) or len(workbooks) != 3
    review_required = exp1_mismatch or (
        not comp_df.empty and bool((comp_df["status"] == "MISMATCH").any())
    )

    verdict = "FAIL" if structural_fail else ("REVIEW REQUIRED" if review_required else "PASS")

    lines.append("")
    lines.append("FINAL VERDICT")
    lines.append("-" * 118)
    lines.append(verdict)
    if verdict == "PASS":
        lines.append("Domain identity evidence is internally aligned for the resolved domains.")
    elif verdict == "REVIEW REQUIRED":
        lines.append("At least one harvest/portable domain identity mismatch requires resolution before Phase-7 freeze.")
    else:
        lines.append("Structural processing failure prevents interpretation.")

    lines.append("")
    lines.append("RESEARCH LOCK STATUS")
    lines.append("-" * 118)
    lines.append("EXP1 domain identity frozen: NO")
    lines.append("Cross-table domain linkage frozen: NO")
    lines.append("Primary target frozen: NO")
    lines.append("Feature set frozen: NO")
    lines.append("Shift family frozen: NO")
    lines.append("Shift severity frozen: NO")
    lines.append("Augmentation frozen: NO")
    lines.append("Model training performed: NO")

    REPORT.write_text("\n".join(lines) + "\n", encoding="utf-8")

    manifest = {
        "phase": "P7.6",
        "archive": str(ARCHIVE),
        "archive_sha256": digest,
        "expected_sha256": EXPECTED_SHA256,
        "expected_sha256_match": digest == EXPECTED_SHA256,
        "verdict": verdict,
        "processing_errors": errors,
        "outputs": {
            "report": str(REPORT),
            "comparison_csv": str(COMPARISON_CSV),
            "raw_labels_csv": str(LABELS_CSV),
            "portable_header_evidence_csv": str(HEADERS_CSV),
        },
        "model_training_performed": False,
        "design_frozen": False,
    }
    MANIFEST.write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    print("\n".join(lines))
    print("")
    print(f"Report: {REPORT}")
    print(f"Comparison: {COMPARISON_CSV}")
    print(f"Header evidence: {HEADERS_CSV}")

    return 1 if structural_fail else 0


if __name__ == "__main__":
    raise SystemExit(main())
