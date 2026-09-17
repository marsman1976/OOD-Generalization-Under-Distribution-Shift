#!/usr/bin/env python3
"""
Phase 7.5b — Raw Temporal Layout Inspection

Diagnostic-only inspection of the original HydroGrowNet Excel workbooks.

Purpose
-------
Resolve the temporal ambiguity found in P7.5 WITHOUT guessing:
1. Show every non-empty raw date-cell candidate from the portable sheet in
   source-row order.
2. Preserve the raw Python/Excel value and type.
3. Show day-first and month-first interpretations side by side.
4. Show surrounding row context for each date candidate.
5. Inspect harvest-sheet title/header regions exactly as stored.
6. Inspect workbook/sheet names for date-like strings, but DO NOT convert those
   strings into authoritative harvest dates.
7. Report Excel merged-cell ranges and workbook date epoch.

No model training.
No data mutation.
No imputation.
No Phase-7 research variable is frozen.
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
import zipfile
from datetime import date, datetime, time
from io import BytesIO
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = ROOT / "data" / "phase7" / "raw"
OUT = ROOT / "results" / "experiment_7" / "temporal_layout_inspection"

PORTABLE_HINT = "water quality parametersportabl"

OUT.mkdir(parents=True, exist_ok=True)


def clean(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass
    return re.sub(r"\s+", " ", str(value)).strip()


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def find_zip() -> Path:
    preferred = RAW_DIR / "all_months_sensory_data.zip"
    if preferred.exists():
        return preferred

    zips = sorted(RAW_DIR.glob("*.zip"))
    if len(zips) == 1:
        return zips[0]
    if not zips:
        raise FileNotFoundError(f"No ZIP archive found in {RAW_DIR}")
    raise RuntimeError(
        "Multiple ZIP archives found and preferred archive is absent:\n"
        + "\n".join(f"  - {p.name}" for p in zips)
    )


def experiment_from_name(name: str) -> str | None:
    m = re.search(r"EXP[.\s_-]*(\d+)", Path(name).name, flags=re.I)
    return f"EXP{int(m.group(1))}" if m else None


def find_portable_sheet(names: list[str]) -> str | None:
    for name in names:
        if name.lower().strip() == PORTABLE_HINT:
            return name
    for name in names:
        if "portabl" in name.lower():
            return name
    return None


def find_harvest_sheet(names: list[str]) -> str | None:
    candidates = [name for name in names if "harvest" in name.lower()]
    return candidates[0] if candidates else None


def value_type(value: Any) -> str:
    if value is None:
        return "None"
    return type(value).__name__


def raw_repr(value: Any) -> str:
    try:
        return repr(value)
    except Exception:
        return str(value)


def numeric_date_parts(text: str) -> tuple[int, int, int] | None:
    s = text.strip().replace("\\", "/").replace("-", "/").replace(".", "/")
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{2,4})", s)
    if not m:
        return None
    a, b, y = map(int, m.groups())
    if y < 100:
        y += 2000
    return a, b, y


def looks_like_date(value: Any) -> bool:
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return True

    s = clean(value)
    if not s:
        return False

    if numeric_date_parts(s) is not None:
        return True

    # Other conservative forms.
    patterns = (
        r"\d{4}[/-]\d{1,2}[/-]\d{1,2}",
        r"\d{1,2}\s+[A-Za-z]{3,9}\s+\d{2,4}",
        r"[A-Za-z]{3,9}\s+\d{1,2},?\s+\d{2,4}",
    )
    return any(re.fullmatch(p, s) for p in patterns)


def parse_dayfirst(value: Any) -> pd.Timestamp:
    if value is None:
        return pd.NaT
    if isinstance(value, pd.Timestamp):
        return value.normalize()
    if isinstance(value, (datetime, date)):
        return pd.Timestamp(value).normalize()

    s = clean(value)
    if not s:
        return pd.NaT
    return pd.to_datetime(s, dayfirst=True, errors="coerce")


def parse_monthfirst(value: Any) -> pd.Timestamp:
    if value is None:
        return pd.NaT
    if isinstance(value, pd.Timestamp):
        return value.normalize()
    if isinstance(value, (datetime, date)):
        return pd.Timestamp(value).normalize()

    s = clean(value)
    if not s:
        return pd.NaT
    return pd.to_datetime(s, dayfirst=False, errors="coerce")


def fmt_date(ts: pd.Timestamp) -> str:
    if pd.isna(ts):
        return ""
    return pd.Timestamp(ts).strftime("%Y-%m-%d")


def ambiguity_status(value: Any) -> str:
    parts = numeric_date_parts(clean(value))
    if parts is None:
        return "not_numeric_slash_date"

    a, b, _ = parts
    if 1 <= a <= 12 and 1 <= b <= 12 and a != b:
        return "ambiguous_day_month"
    if a > 12 and 1 <= b <= 12:
        return "dayfirst_forced_by_first_field_gt_12"
    if b > 12 and 1 <= a <= 12:
        return "monthfirst_forced_by_second_field_gt_12"
    if a == b and 1 <= a <= 12:
        return "same_under_both_orders"
    return "invalid_or_unusual"


def row_context(df: pd.DataFrame, row: int, radius: int = 1) -> str:
    chunks = []
    for rr in range(max(0, row - radius), min(len(df), row + radius + 1)):
        vals = []
        for c in range(df.shape[1]):
            v = clean(df.iat[rr, c])
            if v:
                vals.append(f"C{c}={v!r}")
        chunks.append(f"R{rr}: " + (" | ".join(vals) if vals else "<EMPTY>"))
    return " || ".join(chunks)


def sheet_top_region(df: pd.DataFrame, max_rows: int = 20) -> list[str]:
    lines = []
    for r in range(min(max_rows, len(df))):
        vals = []
        for c in range(df.shape[1]):
            v = clean(df.iat[r, c])
            if v:
                vals.append(f"C{c}={v!r}")
        lines.append(f"R{r}: " + (" | ".join(vals) if vals else "<EMPTY>"))
    return lines


def date_like_substrings(text: str) -> list[str]:
    """
    Extract possible date-like digit groups without interpreting them.
    This deliberately does NOT decide that e.g. '842024' means 8/4/2024.
    """
    if not text:
        return []
    found = []
    for pattern in (
        r"\b\d{1,2}[/-]\d{1,2}[/-]\d{2,4}\b",
        r"\b\d{6,8}\b",
    ):
        found.extend(re.findall(pattern, text))
    return found


def main() -> int:
    archive = find_zip()
    archive_hash = sha256_file(archive)

    portable_records: list[dict[str, Any]] = []
    harvest_candidates: list[dict[str, Any]] = []
    workbook_inventory: list[dict[str, Any]] = []
    report: list[str] = []
    errors: list[str] = []

    report.append("PHASE 7.5b — RAW TEMPORAL LAYOUT INSPECTION")
    report.append("=" * 118)
    report.append(f"Archive: {archive}")
    report.append(f"Archive SHA256: {archive_hash}")
    report.append(
        "Diagnostic only: raw temporal evidence is displayed without choosing a date convention."
    )
    report.append(
        "No model training, data mutation, imputation, or Phase-7 research-variable freeze."
    )
    report.append("")

    with zipfile.ZipFile(archive, "r") as zf:
        excel_members = [
            name
            for name in zf.namelist()
            if not name.endswith("/")
            and Path(name).suffix.lower() in {".xlsx", ".xlsm"}
            and not Path(name).name.startswith("~$")
        ]

        for member in sorted(excel_members):
            experiment = experiment_from_name(member)
            if experiment not in {"EXP1", "EXP2", "EXP3"}:
                continue

            try:
                raw_bytes = zf.read(member)

                # pandas view
                bio = BytesIO(raw_bytes)
                xls = pd.ExcelFile(bio, engine="openpyxl")
                portable = find_portable_sheet(xls.sheet_names)
                harvest = find_harvest_sheet(xls.sheet_names)

                # openpyxl view preserves workbook metadata, merged ranges,
                # formula/value types, and Excel-native date objects.
                wb_formula = load_workbook(
                    BytesIO(raw_bytes),
                    data_only=False,
                    read_only=False,
                )
                wb_values = load_workbook(
                    BytesIO(raw_bytes),
                    data_only=True,
                    read_only=False,
                )

                epoch = str(wb_formula.epoch)

                workbook_inventory.append(
                    {
                        "experiment": experiment,
                        "workbook": Path(member).name,
                        "portable_sheet": portable or "",
                        "harvest_sheet": harvest or "",
                        "excel_epoch": epoch,
                        "workbook_name_digit_candidates": ",".join(
                            date_like_substrings(Path(member).name)
                        ),
                        "harvest_sheet_name_digit_candidates": ",".join(
                            date_like_substrings(harvest or "")
                        ),
                    }
                )

                report.append("=" * 118)
                report.append(
                    f"{experiment} | WORKBOOK: {Path(member).name}"
                )
                report.append(f"Excel epoch: {epoch}")
                report.append(f"Portable sheet: {portable!r}")
                report.append(f"Harvest sheet: {harvest!r}")
                report.append(
                    "Harvest-sheet-name digit candidates (UNINTERPRETED): "
                    + repr(date_like_substrings(harvest or ""))
                )
                report.append("=" * 118)
                report.append("")

                if portable is None:
                    errors.append(f"{experiment}: portable sheet not found")
                else:
                    bio = BytesIO(raw_bytes)
                    portable_df = pd.read_excel(
                        bio,
                        sheet_name=portable,
                        header=None,
                        dtype=object,
                        engine="openpyxl",
                    )

                    ws_formula = wb_formula[portable]
                    ws_values = wb_values[portable]

                    report.append("PORTABLE SHEET MERGED RANGES")
                    report.append("-" * 118)
                    merged = [str(rng) for rng in ws_formula.merged_cells.ranges]
                    report.append(", ".join(merged) if merged else "NONE")
                    report.append("")

                    report.append("RAW DATE-CELL CANDIDATES IN SOURCE-ROW ORDER")
                    report.append("-" * 118)

                    n_candidates = 0
                    for r in range(portable_df.shape[0]):
                        # P7.4/P7.5 established that the portable date field is C0,
                        # but we also inspect all cells for date-like objects/strings.
                        row_candidates = []

                        for c in range(portable_df.shape[1]):
                            value = portable_df.iat[r, c]
                            if not looks_like_date(value):
                                continue

                            row_candidates.append((c, value))

                        for c, value in row_candidates:
                            n_candidates += 1
                            d1 = parse_dayfirst(value)
                            d2 = parse_monthfirst(value)

                            excel_cell = ws_formula.cell(row=r + 1, column=c + 1)
                            excel_value_cell = ws_values.cell(row=r + 1, column=c + 1)

                            rec = {
                                "experiment": experiment,
                                "sheet": portable,
                                "row_zero_based": r,
                                "column_zero_based": c,
                                "excel_coordinate": excel_cell.coordinate,
                                "raw_repr": raw_repr(value),
                                "python_type": value_type(value),
                                "openpyxl_formula_value_repr": raw_repr(excel_cell.value),
                                "openpyxl_data_only_value_repr": raw_repr(excel_value_cell.value),
                                "openpyxl_data_type": excel_cell.data_type,
                                "openpyxl_number_format": excel_cell.number_format,
                                "dayfirst_date": fmt_date(d1),
                                "monthfirst_date": fmt_date(d2),
                                "ambiguity_status": ambiguity_status(value),
                                "interpretations_differ": (
                                    pd.notna(d1)
                                    and pd.notna(d2)
                                    and pd.Timestamp(d1).normalize()
                                    != pd.Timestamp(d2).normalize()
                                ),
                                "row_context": row_context(portable_df, r, radius=1),
                            }
                            portable_records.append(rec)

                            report.append(
                                f"R{r} C{c} ({excel_cell.coordinate}) | "
                                f"raw={raw_repr(value)} | "
                                f"python_type={value_type(value)} | "
                                f"openpyxl_type={excel_cell.data_type!r} | "
                                f"number_format={excel_cell.number_format!r}"
                            )
                            report.append(
                                f"    day-first={fmt_date(d1) or 'UNPARSED'} | "
                                f"month-first={fmt_date(d2) or 'UNPARSED'} | "
                                f"status={ambiguity_status(value)}"
                            )
                            report.append(
                                f"    formula-view={raw_repr(excel_cell.value)} | "
                                f"data-only-view={raw_repr(excel_value_cell.value)}"
                            )
                            report.append(
                                "    context: " + row_context(portable_df, r, radius=1)
                            )

                    if n_candidates == 0:
                        report.append("NONE DETECTED")
                    report.append("")

                    report.append("PORTABLE FIRST 12 ROWS — EXACT NON-EMPTY LAYOUT")
                    report.append("-" * 118)
                    report.extend(sheet_top_region(portable_df, max_rows=12))
                    report.append("")

                if harvest is None:
                    errors.append(f"{experiment}: harvest sheet not found")
                else:
                    bio = BytesIO(raw_bytes)
                    harvest_df = pd.read_excel(
                        bio,
                        sheet_name=harvest,
                        header=None,
                        dtype=object,
                        engine="openpyxl",
                    )

                    ws_formula = wb_formula[harvest]
                    ws_values = wb_values[harvest]

                    report.append("HARVEST SHEET MERGED RANGES")
                    report.append("-" * 118)
                    merged = [str(rng) for rng in ws_formula.merged_cells.ranges]
                    report.append(", ".join(merged) if merged else "NONE")
                    report.append("")

                    report.append("HARVEST TOP 20 ROWS — EXACT NON-EMPTY LAYOUT")
                    report.append("-" * 118)
                    report.extend(sheet_top_region(harvest_df, max_rows=20))
                    report.append("")

                    report.append("HARVEST DATE-LIKE CELLS — FULL SHEET")
                    report.append("-" * 118)
                    found_harvest_date_cell = False

                    for r in range(harvest_df.shape[0]):
                        for c in range(harvest_df.shape[1]):
                            value = harvest_df.iat[r, c]
                            if not looks_like_date(value):
                                continue

                            found_harvest_date_cell = True
                            d1 = parse_dayfirst(value)
                            d2 = parse_monthfirst(value)
                            cell = ws_formula.cell(row=r + 1, column=c + 1)
                            cell_data = ws_values.cell(row=r + 1, column=c + 1)

                            harvest_candidates.append(
                                {
                                    "experiment": experiment,
                                    "sheet": harvest,
                                    "row_zero_based": r,
                                    "column_zero_based": c,
                                    "excel_coordinate": cell.coordinate,
                                    "raw_repr": raw_repr(value),
                                    "python_type": value_type(value),
                                    "openpyxl_formula_value_repr": raw_repr(cell.value),
                                    "openpyxl_data_only_value_repr": raw_repr(cell_data.value),
                                    "openpyxl_data_type": cell.data_type,
                                    "openpyxl_number_format": cell.number_format,
                                    "dayfirst_date": fmt_date(d1),
                                    "monthfirst_date": fmt_date(d2),
                                    "ambiguity_status": ambiguity_status(value),
                                    "row_context": row_context(harvest_df, r, radius=2),
                                }
                            )

                            report.append(
                                f"R{r} C{c} ({cell.coordinate}) | "
                                f"raw={raw_repr(value)} | "
                                f"day-first={fmt_date(d1) or 'UNPARSED'} | "
                                f"month-first={fmt_date(d2) or 'UNPARSED'} | "
                                f"status={ambiguity_status(value)}"
                            )
                            report.append(
                                "    context: " + row_context(harvest_df, r, radius=2)
                            )

                    if not found_harvest_date_cell:
                        report.append("NONE DETECTED")
                    report.append("")

                    report.append(
                        "HARVEST SHEET NAME — POSSIBLE DIGIT GROUPS, NOT INTERPRETED"
                    )
                    report.append("-" * 118)
                    report.append(f"Sheet name: {harvest!r}")
                    digit_groups = date_like_substrings(harvest)
                    report.append(f"Digit groups: {digit_groups!r}")
                    report.append(
                        "Scientific rule: a digit group in a sheet name is NOT accepted "
                        "as a harvest date until its meaning is established from source context."
                    )
                    report.append("")

            except Exception as exc:
                errors.append(
                    f"{experiment} {Path(member).name}: "
                    f"{type(exc).__name__}: {exc}"
                )

    portable_df_out = pd.DataFrame(portable_records)
    harvest_df_out = pd.DataFrame(harvest_candidates)
    workbook_df_out = pd.DataFrame(workbook_inventory)

    report.append("=" * 118)
    report.append("CROSS-EXPERIMENT TEMPORAL EVIDENCE SUMMARY")
    report.append("-" * 118)

    if portable_df_out.empty:
        report.append("Portable date candidates: 0")
    else:
        report.append(f"Portable date candidates: {len(portable_df_out)}")
        report.append(
            "Ambiguous day/month candidates: "
            + str(
                int(
                    (
                        portable_df_out["ambiguity_status"]
                        == "ambiguous_day_month"
                    ).sum()
                )
            )
        )
        report.append(
            "Candidates where day-first/month-first differ: "
            + str(int(portable_df_out["interpretations_differ"].sum()))
        )

        for exp, g in portable_df_out.groupby("experiment", sort=True):
            report.append(f"{exp}:")
            for row in g.sort_values(
                ["row_zero_based", "column_zero_based"]
            ).itertuples(index=False):
                report.append(
                    f"  R{row.row_zero_based} C{row.column_zero_based}: "
                    f"{row.raw_repr} -> "
                    f"day-first={row.dayfirst_date or 'UNPARSED'}, "
                    f"month-first={row.monthfirst_date or 'UNPARSED'}, "
                    f"{row.ambiguity_status}"
                )

    report.append("")
    report.append("HARVEST TEMPORAL EVIDENCE SUMMARY")
    report.append("-" * 118)

    if harvest_df_out.empty:
        report.append(
            "No explicit date-like cell was detected in the harvest sheets."
        )
    else:
        report.append(
            f"Explicit harvest-sheet date-like cells detected: {len(harvest_df_out)}"
        )
        for row in harvest_df_out.itertuples(index=False):
            report.append(
                f"{row.experiment} {row.excel_coordinate}: {row.raw_repr} | "
                f"day-first={row.dayfirst_date or 'UNPARSED'} | "
                f"month-first={row.monthfirst_date or 'UNPARSED'}"
            )

    report.append("")
    report.append("UNRESOLVED ITEMS")
    report.append("-" * 118)
    report.append(
        "1. This diagnostic does not choose day-first vs month-first."
    )
    report.append(
        "2. It does not infer harvest dates solely from worksheet names."
    )
    report.append(
        "3. It does not construct exposure windows or shift severity."
    )
    report.append(
        "4. Any final temporal correction must be justified from the raw evidence "
        "printed above."
    )
    report.append("")

    report.append("VALIDATION")
    report.append("-" * 118)
    report.append(f"Workbook inspection errors: {len(errors)}")
    for err in errors:
        report.append(f"  - {err}")
    report.append(
        f"Experiments inspected: "
        f"{sorted(workbook_df_out['experiment'].unique().tolist()) if not workbook_df_out.empty else []}"
    )
    report.append(
        f"Portable raw date candidates catalogued: {len(portable_df_out)}"
    )
    report.append(
        f"Harvest date-like cells catalogued: {len(harvest_df_out)}"
    )
    report.append("")

    report.append("RESEARCH LOCK STATUS")
    report.append("-" * 118)
    report.append("Primary target frozen: NO")
    report.append("Feature set frozen: NO")
    report.append("Domain definition frozen: NO")
    report.append("Shift family frozen: NO")
    report.append("Shift severity frozen: NO")
    report.append("Augmentation frozen: NO")
    report.append("Model training performed: NO")

    report_path = OUT / "phase7_temporal_layout_report.txt"
    portable_csv = OUT / "phase7_raw_portable_date_cells.csv"
    harvest_csv = OUT / "phase7_raw_harvest_date_cells.csv"
    workbook_csv = OUT / "phase7_temporal_workbook_inventory.csv"
    manifest_path = OUT / "phase7_temporal_layout_manifest.json"

    report_path.write_text("\n".join(report) + "\n", encoding="utf-8")
    portable_df_out.to_csv(portable_csv, index=False)
    harvest_df_out.to_csv(harvest_csv, index=False)
    workbook_df_out.to_csv(workbook_csv, index=False)

    manifest = {
        "phase": "7.5b",
        "purpose": "raw temporal layout inspection",
        "archive": str(archive.relative_to(ROOT)),
        "archive_sha256": archive_hash,
        "experiments_inspected": (
            sorted(workbook_df_out["experiment"].unique().tolist())
            if not workbook_df_out.empty
            else []
        ),
        "portable_date_candidates": int(len(portable_df_out)),
        "harvest_date_like_cells": int(len(harvest_df_out)),
        "inspection_errors": errors,
        "date_convention_selected": False,
        "harvest_dates_frozen": False,
        "model_training_performed": False,
        "phase7_design_frozen": False,
        "outputs": [
            str(report_path.relative_to(ROOT)),
            str(portable_csv.relative_to(ROOT)),
            str(harvest_csv.relative_to(ROOT)),
            str(workbook_csv.relative_to(ROOT)),
        ],
    }
    manifest_path.write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )

    print("=" * 90)
    print("PHASE 7.5b — RAW TEMPORAL LAYOUT INSPECTION")
    print("=" * 90)
    print(f"Experiments inspected: {len(workbook_df_out)}")
    print(f"Portable date candidates: {len(portable_df_out)}")
    print(f"Harvest date-like cells: {len(harvest_df_out)}")
    print(f"Inspection errors: {len(errors)}")
    print(f"Report: {report_path}")
    print("")
    print("NO DATE CONVENTION WAS FROZEN.")
    print("NO MODEL TRAINING WAS PERFORMED.")

    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
