#!/usr/bin/env python3
"""
Phase 7.5d — EXP2 Date-Anomaly Forensic Inspection

Purpose
-------
Inspect the two anomalous EXP2 portable-sheet date entries identified by P7.5c
without silently correcting them.

Known anomaly from P7.5c:
    after reconstructed 2024-05-12,
    source rows contain literal '13/4/2024' and '14/4/2024'.

This diagnostic collects independent context:
1. Exact EXP2 portable rows around the anomaly.
2. Raw Python/openpyxl values, types, formats, formulas, and merged cells.
3. All non-empty neighboring measurement cells.
4. Date-like evidence across ALL EXP2 worksheets.
5. Nearby source-row continuity in the portable sheet.
6. Candidate interpretation comparison:
      H0 = preserve literal 13/04 and 14/04
      H1 = source-entry month typo -> 13/05 and 14/05
   The script reports evidence but does NOT automatically choose H1.
7. Workbook/sheet names and date-like digit groups, left uninterpreted.

No raw data are modified.
No date is corrected.
No ML model is trained.
No Phase-7 scientific design is frozen.
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
import zipfile
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = ROOT / "data" / "phase7" / "raw"
OUT = ROOT / "results" / "experiment_7" / "exp2_date_anomaly"

PORTABLE_HINT = "water quality parametersportabl"
TARGET_RAW_VALUES = {"13/4/2024", "14/4/2024"}

OUT.mkdir(parents=True, exist_ok=True)


def clean(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass
    return re.sub(r"\s+", " ", str(value)).strip()


def raw_repr(value: Any) -> str:
    try:
        return repr(value)
    except Exception:
        return str(value)


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def find_zip() -> Path:
    preferred = RAW_DIR / "all_months_sensory_data.zip"
    if preferred.exists():
        return preferred

    zips = sorted(RAW_DIR.glob("*.zip"))
    if len(zips) == 1:
        return zips[0]
    if not zips:
        raise FileNotFoundError(f"No ZIP archive found in {RAW_DIR}")
    raise RuntimeError(
        "Multiple ZIP archives found and preferred archive is absent:\n"
        + "\n".join(f"  - {p.name}" for p in zips)
    )


def experiment_from_name(name: str) -> str | None:
    m = re.search(r"EXP[.\s_-]*(\d+)", Path(name).name, flags=re.I)
    return f"EXP{int(m.group(1))}" if m else None


def find_portable_sheet(names: list[str]) -> str | None:
    for name in names:
        if name.lower().strip() == PORTABLE_HINT:
            return name
    for name in names:
        if "portabl" in name.lower():
            return name
    return None


def normalize_sep(text: str) -> str:
    return text.strip().replace("\\", "/").replace("-", "/").replace(".", "/")


def parse_literal_ddmmyyyy(value: Any) -> pd.Timestamp:
    s = normalize_sep(clean(value))
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{2,4})", s)
    if not m:
        return pd.NaT
    d, mo, y = map(int, m.groups())
    if y < 100:
        y += 2000
    try:
        return pd.Timestamp(year=y, month=mo, day=d)
    except Exception:
        return pd.NaT


def reconstruct_excel_datetime(value: Any) -> pd.Timestamp:
    if isinstance(value, pd.Timestamp):
        value = value.to_pydatetime()
    if isinstance(value, (datetime, date)):
        try:
            return pd.Timestamp(
                year=value.year,
                month=value.day,
                day=value.month,
            )
        except Exception:
            return pd.NaT
    return pd.NaT


def corrected_sequence_date(value: Any) -> pd.Timestamp:
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return reconstruct_excel_datetime(value)
    return parse_literal_ddmmyyyy(value)


def fmt(ts: pd.Timestamp) -> str:
    if pd.isna(ts):
        return ""
    return pd.Timestamp(ts).strftime("%Y-%m-%d")


def looks_date_like(value: Any) -> bool:
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return True

    s = clean(value)
    if not s:
        return False

    s2 = normalize_sep(s)
    patterns = (
        r"\d{1,2}/\d{1,2}/\d{2,4}",
        r"\d{4}/\d{1,2}/\d{1,2}",
        r"\d{1,2}\s+[A-Za-z]{3,9}\s+\d{2,4}",
        r"[A-Za-z]{3,9}\s+\d{1,2},?\s+\d{2,4}",
    )
    return any(re.fullmatch(p, s2) for p in patterns)


def digit_groups(text: str) -> list[str]:
    if not text:
        return []
    groups = []
    groups.extend(re.findall(r"\b\d{1,2}[/-]\d{1,2}[/-]\d{2,4}\b", text))
    groups.extend(re.findall(r"\b\d{6,8}\b", text))
    return groups


def nonempty_row_cells(df: pd.DataFrame, row: int) -> str:
    vals = []
    for c in range(df.shape[1]):
        v = df.iat[row, c]
        if clean(v):
            vals.append(f"C{c}={raw_repr(v)}")
    return " | ".join(vals) if vals else "<EMPTY>"


def row_context(df: pd.DataFrame, row: int, radius: int = 3) -> list[str]:
    lines = []
    for rr in range(max(0, row - radius), min(df.shape[0], row + radius + 1)):
        marker = ">>" if rr == row else "  "
        lines.append(f"{marker} R{rr}: {nonempty_row_cells(df, rr)}")
    return lines


def column_date_sequence(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for r in range(df.shape[0]):
        raw = df.iat[r, 0]
        ts = corrected_sequence_date(raw)
        if pd.isna(ts):
            continue
        rows.append(
            {
                "row_zero_based": r,
                "raw_repr": raw_repr(raw),
                "raw_type": type(raw).__name__,
                "sequence_date": fmt(ts),
            }
        )

    out = pd.DataFrame(rows)
    if not out.empty:
        out["sequence_ts"] = pd.to_datetime(out["sequence_date"])
        out["step_days"] = out["sequence_ts"].diff().dt.days
        out.drop(columns=["sequence_ts"], inplace=True)
    return out


def candidate_comparison(sequence: pd.DataFrame) -> pd.DataFrame:
    """
    Compare preserving the two literal dates vs treating only their month field
    as a source-entry typo. This is evidence generation, NOT correction.
    """
    if sequence.empty:
        return pd.DataFrame()

    hypotheses = []

    for hypothesis in ("H0_preserve_literal", "H1_month_typo_to_May"):
        temp = sequence.copy()
        dates = pd.to_datetime(temp["sequence_date"])

        if hypothesis == "H1_month_typo_to_May":
            for idx, row in temp.iterrows():
                raw = str(row["raw_repr"]).strip("'\"")
                if raw == "13/4/2024":
                    dates.loc[idx] = pd.Timestamp("2024-05-13")
                elif raw == "14/4/2024":
                    dates.loc[idx] = pd.Timestamp("2024-05-14")

        diffs = dates.diff().dt.days

        hypotheses.append(
            {
                "hypothesis": hypothesis,
                "n_rows": int(len(temp)),
                "backward_steps": int((diffs < 0).sum()),
                "duplicate_dates": int((diffs == 0).sum()),
                "one_day_steps": int((diffs == 1).sum()),
                "gaps_gt_1_day": int((diffs > 1).sum()),
                "largest_abs_step_days": (
                    int(diffs.abs().max())
                    if diffs.notna().any()
                    else 0
                ),
                "first_date": fmt(dates.iloc[0]) if len(dates) else "",
                "last_date": fmt(dates.iloc[-1]) if len(dates) else "",
            }
        )

    return pd.DataFrame(hypotheses)


def main() -> int:
    archive = find_zip()
    archive_hash = sha256_file(archive)

    report: list[str] = []
    anomaly_rows: list[dict[str, Any]] = []
    all_sheet_date_evidence: list[dict[str, Any]] = []
    workbook_inventory: list[dict[str, Any]] = []
    errors: list[str] = []

    report.append("PHASE 7.5d — EXP2 DATE-ANOMALY FORENSIC INSPECTION")
    report.append("=" * 118)
    report.append(f"Archive: {archive}")
    report.append(f"Archive SHA256: {archive_hash}")
    report.append(
        "Purpose: inspect EXP2's final two anomalous date entries without correcting them."
    )
    report.append(
        "No raw mutation, no automatic anomaly correction, no harvest-date inference, "
        "no model training, no Phase-7 design freeze."
    )
    report.append("")

    with zipfile.ZipFile(archive, "r") as zf:
        exp2_members = [
            name
            for name in zf.namelist()
            if not name.endswith("/")
            and Path(name).suffix.lower() in {".xlsx", ".xlsm"}
            and not Path(name).name.startswith("~$")
            and experiment_from_name(name) == "EXP2"
        ]

        if len(exp2_members) != 1:
            errors.append(
                f"Expected exactly one EXP2 workbook; found {len(exp2_members)}: "
                + ", ".join(exp2_members)
            )

        sequence_df = pd.DataFrame()
        comparison_df = pd.DataFrame()

        for member in sorted(exp2_members):
            try:
                raw_bytes = zf.read(member)
                xls = pd.ExcelFile(BytesIO(raw_bytes), engine="openpyxl")
                portable = find_portable_sheet(xls.sheet_names)

                wb_formula = load_workbook(
                    BytesIO(raw_bytes),
                    data_only=False,
                    read_only=False,
                )
                wb_values = load_workbook(
                    BytesIO(raw_bytes),
                    data_only=True,
                    read_only=False,
                )

                workbook_inventory.append(
                    {
                        "experiment": "EXP2",
                        "workbook": Path(member).name,
                        "excel_epoch": str(wb_formula.epoch),
                        "portable_sheet": portable or "",
                        "sheet_names": " || ".join(xls.sheet_names),
                        "workbook_name_digit_groups": ",".join(
                            digit_groups(Path(member).name)
                        ),
                    }
                )

                report.append("WORKBOOK")
                report.append("-" * 118)
                report.append(f"Name: {Path(member).name}")
                report.append(f"Excel epoch: {wb_formula.epoch}")
                report.append(f"Sheets: {xls.sheet_names!r}")
                report.append(
                    f"Workbook-name digit groups, uninterpreted: "
                    f"{digit_groups(Path(member).name)!r}"
                )
                report.append("")

                if portable is None:
                    errors.append("EXP2 portable sheet not found")
                    continue

                portable_df = pd.read_excel(
                    BytesIO(raw_bytes),
                    sheet_name=portable,
                    header=None,
                    dtype=object,
                    engine="openpyxl",
                )
                ws_formula = wb_formula[portable]
                ws_values = wb_values[portable]

                sequence_df = column_date_sequence(portable_df)
                comparison_df = candidate_comparison(sequence_df)

                report.append("PORTABLE SHEET STRUCTURE")
                report.append("-" * 118)
                report.append(f"Sheet: {portable!r}")
                report.append(f"Shape: {portable_df.shape}")
                report.append(
                    "Merged ranges: "
                    + (
                        ", ".join(str(x) for x in ws_formula.merged_cells.ranges)
                        if ws_formula.merged_cells.ranges
                        else "NONE"
                    )
                )
                report.append("")

                report.append("DATE SEQUENCE AROUND FINAL EXP2 ROWS")
                report.append("-" * 118)
                if sequence_df.empty:
                    report.append("No reconstructed date sequence.")
                else:
                    for row in sequence_df.tail(18).itertuples(index=False):
                        step = (
                            "START"
                            if pd.isna(row.step_days)
                            else f"{int(row.step_days):+d} day(s)"
                        )
                        report.append(
                            f"R{row.row_zero_based}: {row.raw_repr} -> "
                            f"{row.sequence_date} | step={step}"
                        )
                report.append("")

                report.append("EXACT ANOMALY-CELL FORENSICS")
                report.append("-" * 118)
                found_targets = 0

                for r in range(portable_df.shape[0]):
                    raw = portable_df.iat[r, 0]
                    normalized = normalize_sep(clean(raw))
                    if normalized not in TARGET_RAW_VALUES:
                        continue

                    found_targets += 1
                    cell = ws_formula.cell(row=r + 1, column=1)
                    cell_value = ws_values.cell(row=r + 1, column=1)

                    rec = {
                        "row_zero_based": r,
                        "excel_coordinate": cell.coordinate,
                        "raw_value": clean(raw),
                        "raw_repr": raw_repr(raw),
                        "python_type": type(raw).__name__,
                        "openpyxl_formula_value_repr": raw_repr(cell.value),
                        "openpyxl_data_only_value_repr": raw_repr(cell_value.value),
                        "openpyxl_data_type": cell.data_type,
                        "openpyxl_number_format": cell.number_format,
                        "literal_ddmmyyyy": fmt(parse_literal_ddmmyyyy(raw)),
                        "full_nonempty_row": nonempty_row_cells(portable_df, r),
                    }
                    anomaly_rows.append(rec)

                    report.append(
                        f"R{r} / {cell.coordinate}: raw={raw_repr(raw)} | "
                        f"python_type={type(raw).__name__} | "
                        f"openpyxl_type={cell.data_type!r} | "
                        f"number_format={cell.number_format!r}"
                    )
                    report.append(
                        f"    formula-view={raw_repr(cell.value)} | "
                        f"data-only-view={raw_repr(cell_value.value)}"
                    )
                    report.append(
                        f"    literal DD/MM/YYYY interpretation="
                        f"{fmt(parse_literal_ddmmyyyy(raw))}"
                    )
                    report.append("    full row:")
                    report.append("      " + nonempty_row_cells(portable_df, r))
                    report.append("    surrounding rows:")
                    for line in row_context(portable_df, r, radius=3):
                        report.append("      " + line)
                    report.append("")

                if found_targets != 2:
                    errors.append(
                        f"Expected 2 target anomaly cells, found {found_targets}"
                    )

                report.append("HYPOTHESIS COMPARISON — DIAGNOSTIC ONLY")
                report.append("-" * 118)
                report.append(
                    "H0: preserve literal 13/4/2024 and 14/4/2024."
                )
                report.append(
                    "H1: test whether changing ONLY their month from 4 to 5 "
                    "would restore source-row continuity."
                )
                report.append(
                    "H1 is NOT applied to the data by this script."
                )
                if comparison_df.empty:
                    report.append("No comparison available.")
                else:
                    for row in comparison_df.itertuples(index=False):
                        report.append(
                            f"{row.hypothesis}: backward={row.backward_steps}, "
                            f"duplicates={row.duplicate_dates}, "
                            f"1-day_steps={row.one_day_steps}, "
                            f"gaps>1day={row.gaps_gt_1_day}, "
                            f"largest_abs_step={row.largest_abs_step_days}, "
                            f"window={row.first_date}..{row.last_date}"
                        )
                report.append("")

                report.append("DATE-LIKE EVIDENCE ACROSS ALL EXP2 SHEETS")
                report.append("-" * 118)

                for sheet in xls.sheet_names:
                    df = pd.read_excel(
                        BytesIO(raw_bytes),
                        sheet_name=sheet,
                        header=None,
                        dtype=object,
                        engine="openpyxl",
                    )
                    ws_f = wb_formula[sheet]
                    ws_v = wb_values[sheet]

                    sheet_found = 0

                    # Sheet-name evidence is recorded separately and not interpreted.
                    groups = digit_groups(sheet)
                    if groups:
                        report.append(
                            f"SHEET NAME {sheet!r}: digit groups={groups!r} "
                            "(UNINTERPRETED)"
                        )

                    for r in range(df.shape[0]):
                        for c in range(df.shape[1]):
                            value = df.iat[r, c]
                            if not looks_date_like(value):
                                continue

                            sheet_found += 1
                            cell = ws_f.cell(row=r + 1, column=c + 1)
                            cell_v = ws_v.cell(row=r + 1, column=c + 1)

                            literal = parse_literal_ddmmyyyy(value)
                            reconstructed = reconstruct_excel_datetime(value)

                            all_sheet_date_evidence.append(
                                {
                                    "sheet": sheet,
                                    "row_zero_based": r,
                                    "column_zero_based": c,
                                    "excel_coordinate": cell.coordinate,
                                    "raw_repr": raw_repr(value),
                                    "python_type": type(value).__name__,
                                    "openpyxl_formula_value_repr": raw_repr(cell.value),
                                    "openpyxl_data_only_value_repr": raw_repr(cell_v.value),
                                    "openpyxl_data_type": cell.data_type,
                                    "openpyxl_number_format": cell.number_format,
                                    "literal_ddmmyyyy_if_string": fmt(literal),
                                    "component_swap_if_excel_datetime": fmt(reconstructed),
                                    "row_context": nonempty_row_cells(df, r),
                                }
                            )

                            report.append(
                                f"{sheet!r} R{r} C{c} ({cell.coordinate}): "
                                f"raw={raw_repr(value)} | type={type(value).__name__} | "
                                f"literal_DMY={fmt(literal) or '-'} | "
                                f"excel_swap={fmt(reconstructed) or '-'}"
                            )

                    if sheet_found == 0:
                        report.append(f"{sheet!r}: no date-like cells detected.")

                report.append("")

            except Exception as exc:
                errors.append(
                    f"{Path(member).name}: {type(exc).__name__}: {exc}"
                )

    anomaly_df = pd.DataFrame(anomaly_rows)
    evidence_df = pd.DataFrame(all_sheet_date_evidence)
    workbook_df = pd.DataFrame(workbook_inventory)

    report.append("EVIDENCE SUMMARY")
    report.append("-" * 118)
    report.append(f"Target anomaly cells recovered: {len(anomaly_df)}")
    report.append(f"Date-like cells catalogued across EXP2 sheets: {len(evidence_df)}")

    if not comparison_df.empty:
        for row in comparison_df.itertuples(index=False):
            report.append(
                f"{row.hypothesis}: backward_steps={row.backward_steps}, "
                f"one_day_steps={row.one_day_steps}, "
                f"largest_abs_step={row.largest_abs_step_days}"
            )

    report.append("")
    report.append("INTERPRETATION RULE")
    report.append("-" * 118)
    report.append(
        "This script does NOT declare 13/4 and 14/4 to be errors merely because "
        "13/5 and 14/5 would produce a smoother sequence."
    )
    report.append(
        "A correction should be recorded only if the combined raw workbook context "
        "provides sufficient provenance/evidence for that decision."
    )
    report.append("")

    report.append("VALIDATION")
    report.append("-" * 118)
    report.append(f"Processing errors: {len(errors)}")
    for err in errors:
        report.append(f"  - {err}")
    report.append(f"EXP2 workbook count: {len(workbook_df)}")
    report.append(f"Anomaly cells found: {len(anomaly_df)}")
    report.append("")

    report.append("RESEARCH LOCK STATUS")
    report.append("-" * 118)
    report.append("EXP2 anomaly corrected: NO")
    report.append("Environmental chronology frozen: NO")
    report.append("Harvest dates frozen: NO")
    report.append("Primary target frozen: NO")
    report.append("Feature set frozen: NO")
    report.append("Domain definition frozen: NO")
    report.append("Shift family frozen: NO")
    report.append("Shift severity frozen: NO")
    report.append("Augmentation frozen: NO")
    report.append("Model training performed: NO")

    report_path = OUT / "phase7_exp2_date_anomaly_report.txt"
    anomaly_csv = OUT / "phase7_exp2_anomaly_cells.csv"
    evidence_csv = OUT / "phase7_exp2_all_sheet_date_evidence.csv"
    sequence_csv = OUT / "phase7_exp2_portable_date_sequence.csv"
    comparison_csv = OUT / "phase7_exp2_date_hypothesis_comparison.csv"
    workbook_csv = OUT / "phase7_exp2_workbook_inventory.csv"
    manifest_path = OUT / "phase7_exp2_date_anomaly_manifest.json"

    report_path.write_text("\n".join(report) + "\n", encoding="utf-8")
    anomaly_df.to_csv(anomaly_csv, index=False)
    evidence_df.to_csv(evidence_csv, index=False)
    sequence_df.to_csv(sequence_csv, index=False)
    comparison_df.to_csv(comparison_csv, index=False)
    workbook_df.to_csv(workbook_csv, index=False)

    manifest = {
        "phase": "7.5d",
        "purpose": "EXP2 date-anomaly forensic inspection",
        "archive": str(archive.relative_to(ROOT)),
        "archive_sha256": archive_hash,
        "target_literal_values": sorted(TARGET_RAW_VALUES),
        "target_anomaly_cells_found": int(len(anomaly_df)),
        "date_like_cells_across_exp2_sheets": int(len(evidence_df)),
        "automatic_correction_applied": False,
        "harvest_date_inferred": False,
        "model_training_performed": False,
        "phase7_design_frozen": False,
        "processing_errors": errors,
        "outputs": [
            str(report_path.relative_to(ROOT)),
            str(anomaly_csv.relative_to(ROOT)),
            str(evidence_csv.relative_to(ROOT)),
            str(sequence_csv.relative_to(ROOT)),
            str(comparison_csv.relative_to(ROOT)),
            str(workbook_csv.relative_to(ROOT)),
        ],
    }
    manifest_path.write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )

    print("=" * 90)
    print("PHASE 7.5d — EXP2 DATE-ANOMALY FORENSIC INSPECTION")
    print("=" * 90)
    print(f"EXP2 workbooks inspected: {len(workbook_df)}")
    print(f"Target anomaly cells found: {len(anomaly_df)}")
    print(f"Date-like cells across EXP2 sheets: {len(evidence_df)}")
    print(f"Processing errors: {len(errors)}")
    print(f"Report: {report_path}")
    print("")
    print("NO AUTOMATIC DATE CORRECTION WAS APPLIED.")
    print("NO MODEL TRAINING WAS PERFORMED.")

    return 0 if len(errors) == 0 and len(anomaly_df) == 2 else 1


if __name__ == "__main__":
    sys.exit(main())
