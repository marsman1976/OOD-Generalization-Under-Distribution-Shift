#!/usr/bin/env python3
"""
Phase 7.6b — EXP1 Raw Domain-Layout Forensic Inspection

Purpose
-------
Resolve the EXP1 cross-table identity mismatch found in P7.6 by inspecting the
raw header hierarchy and domain labels in the EXP1 portable and harvest sheets.

This script is diagnostic only. It does NOT:
- modify the source archive/workbooks
- change any parser or domain mapping
- infer a final treatment/replicate orientation
- train models
- freeze the Phase-7 design
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
ARCHIVE = ROOT / "data" / "phase7" / "raw" / "all_months_sensory_data.zip"
OUT = ROOT / "results" / "experiment_7" / "exp1_domain_layout_inspection"

REPORT = OUT / "phase7_exp1_domain_layout_report.txt"
PORTABLE_CSV = OUT / "phase7_exp1_portable_header_cells.csv"
MERGES_CSV = OUT / "phase7_exp1_portable_merged_ranges.csv"
BLOCKS_CSV = OUT / "phase7_exp1_portable_measurement_blocks.csv"
HARVEST_CSV = OUT / "phase7_exp1_harvest_domain_evidence.csv"
MAPPING_CSV = OUT / "phase7_exp1_candidate_mapping_comparison.csv"
MANIFEST = OUT / "phase7_exp1_domain_layout_manifest.json"

EXPECTED_SHA256 = "daa6d1cc259234d4b1cb8ba9901b2f7e35975a444154d5b8dd0a58a394fdbf14"


def clean(v: Any) -> str:
    if v is None:
        return ""
    try:
        if pd.isna(v):
            return ""
    except Exception:
        pass
    return re.sub(r"\s+", " ", str(v)).strip()


def compact(v: Any) -> str:
    return re.sub(r"[^A-Z0-9]+", "", clean(v).upper())


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def exp_number(name: str) -> int | None:
    m = re.search(r"EXP\.?\s*(\d+)", Path(name).name, flags=re.I)
    return int(m.group(1)) if m else None


def choose_sheet(names: list[str], hint: str) -> str | None:
    for name in names:
        if hint.lower() in name.lower():
            return name
    return None


def explicit_domain(v: Any) -> str | None:
    m = re.fullmatch(r"R(\d+)T(\d+)", compact(v))
    if not m:
        return None
    return f"R{int(m.group(1))}-T{int(m.group(2))}"


def number_after(prefix: str, v: Any) -> int | None:
    s = clean(v).upper()
    if prefix == "R":
        pats = [
            r"\bR\s*[-_:]?\s*(\d+)\b",
            r"\bREPLICATE\s*[-_:]?\s*(\d+)\b",
        ]
    else:
        pats = [
            r"\bT\s*[-_:]?\s*(\d+)\b",
            r"\bTREATMENT\s*[-_:]?\s*(\d+)\b",
        ]
    for p in pats:
        m = re.search(p, s)
        if m:
            return int(m.group(1))
    return None


def value_covering(ws, row: int, col: int) -> tuple[Any, str]:
    """Get direct value or merged anchor value covering a cell."""
    cell = ws.cell(row=row, column=col)
    if cell.value is not None:
        return cell.value, cell.coordinate

    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            anchor = ws.cell(rng.min_row, rng.min_col)
            return anchor.value, f"{rng} -> {anchor.coordinate}"
    return None, ""


def header_snapshot(ws, max_rows: int = 5) -> list[dict]:
    rows = []
    for r in range(1, min(max_rows, ws.max_row) + 1):
        for c in range(1, ws.max_column + 1):
            direct = ws.cell(r, c).value
            covered, source = value_covering(ws, r, c)
            if clean(direct) or clean(covered):
                rows.append({
                    "row": r,
                    "column": c,
                    "coordinate": ws.cell(r, c).coordinate,
                    "column_letter": get_column_letter(c),
                    "direct_value": clean(direct),
                    "effective_header_value": clean(covered),
                    "effective_source": source,
                    "replicate_number": number_after("R", covered),
                    "treatment_number": number_after("T", covered),
                    "explicit_domain": explicit_domain(covered) or "",
                })
    return rows


def merged_snapshot(ws, max_rows: int = 5) -> list[dict]:
    rows = []
    for rng in ws.merged_cells.ranges:
        if rng.min_row > max_rows:
            continue
        anchor = ws.cell(rng.min_row, rng.min_col)
        rows.append({
            "range": str(rng),
            "min_row": rng.min_row,
            "max_row": rng.max_row,
            "min_col": rng.min_col,
            "max_col": rng.max_col,
            "width": rng.max_col - rng.min_col + 1,
            "height": rng.max_row - rng.min_row + 1,
            "anchor": anchor.coordinate,
            "value": clean(anchor.value),
            "replicate_number": number_after("R", anchor.value),
            "treatment_number": number_after("T", anchor.value),
            "explicit_domain": explicit_domain(anchor.value) or "",
        })
    return sorted(rows, key=lambda x: (x["min_row"], x["min_col"]))


def find_measurement_blocks(ws, max_rows: int = 5) -> list[dict]:
    """
    Identify pH starts and inspect exactly which header labels geometrically cover
    each four-column pH/EC/TDS/water-temperature block.
    """
    blocks = []
    seen = set()

    for r in range(1, min(max_rows, ws.max_row) + 1):
        for c in range(1, ws.max_column + 1):
            if clean(ws.cell(r, c).value).lower() != "ph":
                continue
            if c in seen:
                continue
            seen.add(c)

            end = min(c + 3, ws.max_column)
            variables = [clean(ws.cell(r, cc).value) for cc in range(c, end + 1)]

            context = []
            reps = []
            treats = []

            for hr in range(1, r + 1):
                vals = []
                for cc in range(c, end + 1):
                    v, src = value_covering(ws, hr, cc)
                    if clean(v):
                        vals.append((clean(v), src))
                        rv = number_after("R", v)
                        tv = number_after("T", v)
                        if rv is not None:
                            reps.append(rv)
                        if tv is not None:
                            treats.append(tv)
                for v, src in vals:
                    context.append(f"row{hr}:{src}:{v}")

            reps_u = sorted(set(reps))
            treats_u = sorted(set(treats))

            direct_mapping = ""
            if len(reps_u) == 1 and len(treats_u) == 1:
                direct_mapping = f"R{reps_u[0]}-T{treats_u[0]}"

            blocks.append({
                "ph_header_row": r,
                "start_col": c,
                "end_col": end,
                "column_range": f"{get_column_letter(c)}:{get_column_letter(end)}",
                "variables": " | ".join(variables),
                "header_context": " || ".join(sorted(set(context))),
                "replicate_candidates": ",".join(map(str, reps_u)),
                "treatment_candidates": ",".join(map(str, treats_u)),
                "direct_mapping": direct_mapping,
            })

    return blocks


def harvest_evidence(ws) -> list[dict]:
    rows = []
    for row in ws.iter_rows():
        for cell in row:
            d = explicit_domain(cell.value)
            if d:
                m = re.fullmatch(r"R(\d+)-T(\d+)", d)
                rows.append({
                    "cell": cell.coordinate,
                    "raw_value": clean(cell.value),
                    "domain": d,
                    "replicate": int(m.group(1)),
                    "treatment": int(m.group(2)),
                })
    return rows


def transpose_domain(d: str) -> str:
    m = re.fullmatch(r"R(\d+)-T(\d+)", d)
    if not m:
        return d
    return f"R{int(m.group(2))}-T{int(m.group(1))}"


def mapping_comparison(blocks: list[dict], harvest_domains: set[str]) -> list[dict]:
    """
    Compare two hypotheses only:
      H0 direct portable interpretation Rr-Tt
      H1 swap roles Rr-Tt -> Rt-Tr

    This is diagnostic evidence, not an automatic mapping decision.
    """
    direct = [b["direct_mapping"] for b in blocks if b["direct_mapping"]]
    direct_set = set(direct)
    transposed_set = {transpose_domain(d) for d in direct_set}

    hypotheses = [
        ("H0_direct_portable_labels", direct_set),
        ("H1_swap_replicate_treatment_roles", transposed_set),
    ]

    out = []
    for name, candidate in hypotheses:
        out.append({
            "hypothesis": name,
            "candidate_domains": " | ".join(sorted(candidate)),
            "harvest_domains": " | ".join(sorted(harvest_domains)),
            "intersection_count": len(candidate & harvest_domains),
            "candidate_only": " | ".join(sorted(candidate - harvest_domains)),
            "harvest_only": " | ".join(sorted(harvest_domains - candidate)),
            "exact_set_match": bool(candidate) and candidate == harvest_domains,
        })
    return out


def format_grid(ws, rows: int = 5) -> list[str]:
    """Human-readable top-of-sheet layout preserving column positions."""
    lines = []
    max_r = min(rows, ws.max_row)
    for r in range(1, max_r + 1):
        vals = []
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if clean(v):
                vals.append(f"{get_column_letter(c)}{r}={clean(v)!r}")
        lines.append(f"Row {r}: " + (" | ".join(vals) if vals else "<empty>"))
    return lines


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    errors: list[str] = []

    report = [
        "PHASE 7.6b — EXP1 RAW DOMAIN-LAYOUT FORENSIC INSPECTION",
        "=" * 118,
        f"Archive: {ARCHIVE}",
        "Purpose: resolve whether EXP1 portable replicate/treatment roles were interpreted in the correct orientation.",
        "No raw mutation, no parser change, no automatic mapping change, no model training, no Phase-7 design freeze.",
        "",
    ]

    if not ARCHIVE.exists():
        print(f"ERROR: archive not found: {ARCHIVE}")
        return 1

    digest = sha256_file(ARCHIVE)
    report.extend([
        f"Archive SHA256: {digest}",
        f"Expected SHA256 match: {'YES' if digest == EXPECTED_SHA256 else 'NO'}",
        "",
    ])

    exp1_member = None
    with zipfile.ZipFile(ARCHIVE, "r") as zf:
        members = [
            n for n in zf.namelist()
            if n.lower().endswith(".xlsx") and exp_number(n) == 1
        ]
        if len(members) != 1:
            errors.append(f"Expected exactly one EXP1 workbook, found {len(members)}: {members}")
        elif members:
            exp1_member = members[0]

        if exp1_member is None:
            REPORT.write_text("\n".join(report + errors), encoding="utf-8")
            return 1

        raw = zf.read(exp1_member)

    try:
        wb = openpyxl.load_workbook(BytesIO(raw), data_only=False)
    except Exception as exc:
        print(f"ERROR loading EXP1 workbook: {exc}")
        return 1

    portable_name = choose_sheet(wb.sheetnames, "portabl")
    harvest_name = choose_sheet(wb.sheetnames, "harvest")

    report.extend([
        "WORKBOOK / SHEET DISCOVERY",
        "-" * 118,
        f"Workbook: {Path(exp1_member).name!r}",
        f"Sheets: {wb.sheetnames!r}",
        f"Portable sheet: {portable_name!r}",
        f"Harvest sheet: {harvest_name!r}",
        "",
    ])

    if portable_name is None or harvest_name is None:
        errors.append("Required portable or harvest sheet not found.")
        REPORT.write_text("\n".join(report + errors), encoding="utf-8")
        return 1

    pws = wb[portable_name]
    hws = wb[harvest_name]

    header_rows = header_snapshot(pws)
    merge_rows = merged_snapshot(pws)
    blocks = find_measurement_blocks(pws)
    harvest_rows = harvest_evidence(hws)

    pd.DataFrame(header_rows).to_csv(PORTABLE_CSV, index=False)
    pd.DataFrame(merge_rows).to_csv(MERGES_CSV, index=False)
    pd.DataFrame(blocks).to_csv(BLOCKS_CSV, index=False)
    pd.DataFrame(harvest_rows).to_csv(HARVEST_CSV, index=False)

    harvest_domains = {r["domain"] for r in harvest_rows}
    mapping_rows = mapping_comparison(blocks, harvest_domains)
    pd.DataFrame(mapping_rows).to_csv(MAPPING_CSV, index=False)

    report.extend([
        "RAW PORTABLE TOP-HEADER GRID",
        "-" * 118,
    ])
    report.extend(format_grid(pws, rows=5))
    report.append("")

    report.extend([
        "PORTABLE MERGED-RANGE HIERARCHY",
        "-" * 118,
    ])
    for x in merge_rows:
        report.append(
            f"{x['range']}: value={x['value']!r}, "
            f"R={x['replicate_number']}, T={x['treatment_number']}, "
            f"explicit_domain={x['explicit_domain'] or '-'}"
        )
    report.append("")

    report.extend([
        "PORTABLE MEASUREMENT BLOCKS",
        "-" * 118,
    ])
    for i, b in enumerate(blocks, start=1):
        report.append(
            f"Block {i}: cols={b['column_range']}, variables={b['variables']}, "
            f"R_candidates={b['replicate_candidates'] or '-'}, "
            f"T_candidates={b['treatment_candidates'] or '-'}, "
            f"direct_mapping={b['direct_mapping'] or 'UNRESOLVED'}"
        )
        report.append(f"  context: {b['header_context']}")
    report.append("")

    report.extend([
        "HARVEST EXPLICIT DOMAIN LABELS",
        "-" * 118,
        f"Unique harvest domains ({len(harvest_domains)}): {sorted(harvest_domains)}",
    ])
    for x in harvest_rows:
        report.append(f"{x['cell']}: {x['raw_value']!r} -> {x['domain']}")
    report.append("")

    report.extend([
        "CANDIDATE ORIENTATION COMPARISON",
        "-" * 118,
        "H0 preserves the current portable R/T interpretation.",
        "H1 swaps the semantic roles of replicate and treatment in the portable mapping.",
        "Neither hypothesis is automatically applied.",
    ])
    for x in mapping_rows:
        report.append(
            f"{x['hypothesis']}: exact_set_match={x['exact_set_match']}, "
            f"intersection={x['intersection_count']}, "
            f"candidate_only=[{x['candidate_only']}], harvest_only=[{x['harvest_only']}]"
        )
        report.append(f"  candidate domains: {x['candidate_domains']}")
    report.append("")

    h0 = next((x for x in mapping_rows if x["hypothesis"].startswith("H0")), None)
    h1 = next((x for x in mapping_rows if x["hypothesis"].startswith("H1")), None)

    report.extend([
        "FORENSIC INTERPRETATION",
        "-" * 118,
    ])

    if h1 and h1["exact_set_match"] and h0 and not h0["exact_set_match"]:
        report.append(
            "The role-swapped portable candidate reproduces the complete explicit harvest domain set, "
            "whereas the current direct portable interpretation does not."
        )
        report.append(
            "This is strong structural evidence that EXP1's portable header hierarchy was previously "
            "interpreted with replicate and treatment roles transposed."
        )
        report.append(
            "This diagnostic still does not change the parser; a subsequent correction step must preserve "
            "the raw labels and document the mapping provenance."
        )
        forensic_status = "STRONG_EVIDENCE_FOR_ROLE_SWAP"
    elif h0 and h0["exact_set_match"]:
        report.append(
            "The direct portable interpretation already reproduces the harvest domain set."
        )
        forensic_status = "DIRECT_MAPPING_SUPPORTED"
    else:
        report.append(
            "Neither simple orientation reproduces the complete harvest domain set. "
            "Further source-level investigation is required."
        )
        forensic_status = "UNRESOLVED"

    report.extend([
        "",
        "VALIDATION",
        "-" * 118,
        f"Processing errors: {len(errors)}",
        f"Portable header evidence rows: {len(header_rows)}",
        f"Portable merged ranges inspected: {len(merge_rows)}",
        f"Portable pH measurement blocks detected: {len(blocks)}",
        f"Harvest explicit domain-label occurrences: {len(harvest_rows)}",
        f"Forensic status: {forensic_status}",
    ])
    for e in errors:
        report.append(f"ERROR: {e}")

    report.extend([
        "",
        "RESEARCH LOCK STATUS",
        "-" * 118,
        "EXP1 parser changed: NO",
        "EXP1 domain identity frozen: NO",
        "Cross-table domain linkage frozen: NO",
        "Primary target frozen: NO",
        "Feature set frozen: NO",
        "Shift family frozen: NO",
        "Shift severity frozen: NO",
        "Augmentation frozen: NO",
        "Model training performed: NO",
    ])

    REPORT.write_text("\n".join(report) + "\n", encoding="utf-8")

    manifest = {
        "phase": "P7.6b",
        "archive_sha256": digest,
        "expected_sha256_match": digest == EXPECTED_SHA256,
        "workbook": Path(exp1_member).name,
        "portable_sheet": portable_name,
        "harvest_sheet": harvest_name,
        "forensic_status": forensic_status,
        "processing_errors": errors,
        "parser_changed": False,
        "model_training_performed": False,
        "phase7_design_frozen": False,
        "outputs": {
            "report": str(REPORT),
            "portable_header_cells": str(PORTABLE_CSV),
            "portable_merged_ranges": str(MERGES_CSV),
            "portable_measurement_blocks": str(BLOCKS_CSV),
            "harvest_domain_evidence": str(HARVEST_CSV),
            "candidate_mapping_comparison": str(MAPPING_CSV),
        },
    }
    MANIFEST.write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    print("\n".join(report))
    print("")
    print(f"Report: {REPORT}")
    print("NO PARSER OR DOMAIN MAPPING WAS CHANGED.")
    print("NO MODEL TRAINING WAS PERFORMED.")

    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
