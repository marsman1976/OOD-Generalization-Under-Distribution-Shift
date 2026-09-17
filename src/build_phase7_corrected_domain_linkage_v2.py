#!/usr/bin/env python3
"""
Phase 7.6c-v2 — Corrected Cross-Table Environmental Domain Linkage

Purpose
-------
Reconstruct the portable environmental domain identities using the documented
P7.6b EXP1 role-swap correction and verify identity-level agreement with the
harvest domains across EXP1, EXP2, and EXP3.

Correction rule
---------------
EXP1 only:
    corrected_replicate = raw_portable_treatment
    corrected_treatment = raw_portable_replicate

EXP2 and EXP3:
    corrected_replicate = raw_portable_replicate
    corrected_treatment = raw_portable_treatment

Scientific safeguards
---------------------
- Raw workbook labels are preserved in output columns.
- Raw ZIP/workbooks are NEVER modified.
- The EXP1 correction is explicit and provenance-labelled.
- No target, feature set, shift definition, severity metric, augmentation,
  or model configuration is frozen.
- No model is trained.
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
ARCHIVE = ROOT / "data" / "phase7" / "raw" / "all_months_sensory_data.zip"
OUT = ROOT / "results" / "experiment_7" / "corrected_domain_linkage"

REPORT = OUT / "phase7_corrected_domain_linkage_report.txt"
BLOCKS_CSV = OUT / "phase7_corrected_portable_blocks.csv"
LINKAGE_CSV = OUT / "phase7_corrected_domain_linkage.csv"
HARVEST_CSV = OUT / "phase7_harvest_domain_keys.csv"
MANIFEST = OUT / "phase7_corrected_domain_linkage_manifest.json"

EXPECTED_SHA256 = "daa6d1cc259234d4b1cb8ba9901b2f7e35975a444154d5b8dd0a58a394fdbf14"

CORRECTION_ID = "P7.6b_EXP1_ROLE_SWAP"
CORRECTION_RULE = (
    "EXP1 only: corrected_replicate=raw_portable_treatment; "
    "corrected_treatment=raw_portable_replicate"
)


def clean(v: Any) -> str:
    if v is None:
        return ""
    try:
        if pd.isna(v):
            return ""
    except Exception:
        pass
    return re.sub(r"\s+", " ", str(v)).strip()


def compact(v: Any) -> str:
    return re.sub(r"[^A-Z0-9]+", "", clean(v).upper())


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def experiment_number(name: str) -> int | None:
    m = re.search(r"EXP\.?\s*(\d+)", Path(name).name, flags=re.I)
    return int(m.group(1)) if m else None


def choose_sheet(names: list[str], hint: str) -> str | None:
    for name in names:
        if hint.lower() in name.lower():
            return name
    return None


def explicit_domain(v: Any) -> str | None:
    m = re.fullmatch(r"R(\d+)T(\d+)", compact(v))
    if not m:
        return None
    return f"R{int(m.group(1))}-T{int(m.group(2))}"


def parse_portable_label(v: Any) -> tuple[int, int] | None:
    """
    Parse labels such as:
      Replicate 1 T1
      Replicate 2 T2
    Returns (raw_replicate, raw_treatment).
    """
    s = clean(v).upper()
    if not s:
        return None

    r = re.search(r"\bREPLICATE\s*[-_:]?\s*(\d+)\b", s)
    t = re.search(r"\bT\s*[-_:]?\s*(\d+)\b", s)

    if r and t:
        return int(r.group(1)), int(t.group(1))
    return None


def corrected_identity(
    exp: int,
    raw_replicate: int,
    raw_treatment: int,
) -> tuple[int, int, bool, str]:
    if exp == 1:
        return (
            raw_treatment,
            raw_replicate,
            True,
            CORRECTION_ID,
        )
    return (
        raw_replicate,
        raw_treatment,
        False,
        "NONE",
    )


def find_exp_workbooks(zf: zipfile.ZipFile) -> dict[int, str]:
    found: dict[int, str] = {}
    for member in zf.namelist():
        if not member.lower().endswith(".xlsx"):
            continue
        exp = experiment_number(member)
        if exp in (1, 2, 3):
            found[exp] = member
    return found


def harvest_domains(ws, exp: int) -> list[dict]:
    rows: list[dict] = []
    seen: set[str] = set()

    for row in ws.iter_rows():
        for cell in row:
            d = explicit_domain(cell.value)
            if not d or d in seen:
                continue

            seen.add(d)
            m = re.fullmatch(r"R(\d+)-T(\d+)", d)
            rows.append({
                "experiment": f"EXP{exp}",
                "domain": d,
                "replicate": int(m.group(1)),
                "treatment": int(m.group(2)),
                "source_sheet": ws.title,
                "source_cell": cell.coordinate,
                "raw_label": clean(cell.value),
            })

    return sorted(rows, key=lambda x: (x["replicate"], x["treatment"]))


def portable_blocks(ws, exp: int) -> list[dict]:
    """
    Recover every 4-column pH/EC/TDS/water-temperature block from the first
    few header rows and preserve both raw and corrected identities.
    """
    blocks: list[dict] = []

    # In these workbooks the domain labels are merged in row 2 and variables
    # are in row 3. We still search rows 1..5 conservatively.
    for r in range(1, min(ws.max_row, 5) + 1):
        for c in range(1, ws.max_column + 1):
            header_token = re.sub(
                r"[^a-z0-9]+",
                "",
                clean(ws.cell(r, c).value).lower(),
            )
            if header_token != "ph":
                continue

            # Search upward at this column for the nearest portable domain label.
            parsed = None
            label_cell = None
            raw_label = ""

            for hr in range(r - 1, 0, -1):
                # Direct value.
                v = ws.cell(hr, c).value
                p = parse_portable_label(v)
                if p:
                    parsed = p
                    label_cell = ws.cell(hr, c).coordinate
                    raw_label = clean(v)
                    break

                # Merged range covering this coordinate.
                for rng in ws.merged_cells.ranges:
                    if (
                        rng.min_row <= hr <= rng.max_row
                        and rng.min_col <= c <= rng.max_col
                    ):
                        anchor = ws.cell(rng.min_row, rng.min_col)
                        p = parse_portable_label(anchor.value)
                        if p:
                            parsed = p
                            label_cell = f"{rng} -> {anchor.coordinate}"
                            raw_label = clean(anchor.value)
                            break
                if parsed:
                    break

            if not parsed:
                continue

            raw_r, raw_t = parsed
            corr_r, corr_t, corrected, provenance = corrected_identity(
                exp, raw_r, raw_t
            )

            # Determine time-group label from row 1 / merged anchor.
            time_label = ""
            for rng in ws.merged_cells.ranges:
                if rng.min_row <= 1 <= rng.max_row and rng.min_col <= c <= rng.max_col:
                    time_label = clean(ws.cell(rng.min_row, rng.min_col).value)
                    break
            if not time_label:
                time_label = clean(ws.cell(1, c).value)

            variables = [
                clean(ws.cell(r, cc).value)
                for cc in range(c, min(c + 3, ws.max_column) + 1)
            ]

            blocks.append({
                "experiment": f"EXP{exp}",
                "source_sheet": ws.title,
                "time_label": time_label,
                "block_start_cell": ws.cell(r, c).coordinate,
                "block_columns": (
                    f"{get_column_letter(c)}:"
                    f"{get_column_letter(min(c + 3, ws.max_column))}"
                ),
                "variables": " | ".join(variables),
                "raw_label": raw_label,
                "raw_label_source": label_cell or "",
                "raw_portable_replicate": raw_r,
                "raw_portable_treatment": raw_t,
                "raw_portable_domain": f"R{raw_r}-T{raw_t}",
                "correction_applied": corrected,
                "correction_provenance": provenance,
                "corrected_replicate": corr_r,
                "corrected_treatment": corr_t,
                "corrected_domain": f"R{corr_r}-T{corr_t}",
            })

    return blocks


def unique_environment_domains(blocks: list[dict], exp: int) -> set[str]:
    return {
        b["corrected_domain"]
        for b in blocks
        if b["experiment"] == f"EXP{exp}"
    }


def unique_raw_domains(blocks: list[dict], exp: int) -> set[str]:
    return {
        b["raw_portable_domain"]
        for b in blocks
        if b["experiment"] == f"EXP{exp}"
    }


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)

    report: list[str] = [
        "PHASE 7.6c-v2 — CORRECTED CROSS-TABLE ENVIRONMENTAL DOMAIN LINKAGE",
        "=" * 118,
        f"Archive: {ARCHIVE}",
        "Purpose: apply the documented EXP1 semantic role-swap to the analytical linkage only,",
        "while preserving raw portable labels and verifying identity-level harvest/environment agreement.",
        "No raw workbook mutation. No model training. No Phase-7 design freeze.",
        "P7.6d parser fix: measurement-header matching ignores punctuation/case/spacing,",
        "so source header 'pH.' is recognized as pH without altering the workbook.",
        "",
    ]
    errors: list[str] = []

    if not ARCHIVE.exists():
        print(f"ERROR: archive not found: {ARCHIVE}")
        return 1

    digest = sha256_file(ARCHIVE)
    report.extend([
        f"Archive SHA256: {digest}",
        f"Expected SHA256 match: {'YES' if digest == EXPECTED_SHA256 else 'NO'}",
        "",
        "DOCUMENTED CORRECTION",
        "-" * 118,
        f"Correction ID: {CORRECTION_ID}",
        f"Rule: {CORRECTION_RULE}",
        "EXP2/EXP3 rule: corrected identity = raw portable identity.",
        "Raw labels remain preserved in all outputs.",
        "",
    ])

    all_blocks: list[dict] = []
    all_harvest: list[dict] = []
    workbook_names: dict[int, str] = {}

    with zipfile.ZipFile(ARCHIVE, "r") as zf:
        workbooks = find_exp_workbooks(zf)

        for exp in (1, 2, 3):
            member = workbooks.get(exp)
            if member is None:
                errors.append(f"EXP{exp}: workbook not found")
                continue

            workbook_names[exp] = Path(member).name

            try:
                raw = zf.read(member)
                wb = openpyxl.load_workbook(BytesIO(raw), data_only=False)
            except Exception as exc:
                errors.append(
                    f"EXP{exp}: workbook load failed: {type(exc).__name__}: {exc}"
                )
                continue

            portable_name = choose_sheet(wb.sheetnames, "portabl")
            harvest_name = choose_sheet(wb.sheetnames, "harvest")

            if portable_name is None:
                errors.append(f"EXP{exp}: portable sheet not found")
                continue
            if harvest_name is None:
                errors.append(f"EXP{exp}: harvest sheet not found")
                continue

            pblocks = portable_blocks(wb[portable_name], exp)
            hdomains = harvest_domains(wb[harvest_name], exp)

            all_blocks.extend(pblocks)
            all_harvest.extend(hdomains)

            report.extend([
                f"EXP{exp} SOURCE DISCOVERY",
                "-" * 118,
                f"Workbook: {Path(member).name!r}",
                f"Portable sheet: {portable_name!r}",
                f"Harvest sheet: {harvest_name!r}",
                f"Portable measurement blocks recovered: {len(pblocks)}",
                f"Harvest unique explicit domains recovered: {len(hdomains)}",
                "",
            ])

    blocks_df = pd.DataFrame(all_blocks)
    harvest_df = pd.DataFrame(all_harvest)

    blocks_df.to_csv(BLOCKS_CSV, index=False)
    harvest_df.to_csv(HARVEST_CSV, index=False)

    linkage_rows: list[dict] = []
    total_match = 0
    total_mismatch = 0

    report.extend([
        "IDENTITY-LEVEL LINKAGE VERIFICATION",
        "=" * 118,
    ])

    for exp in (1, 2, 3):
        hset = {
            x["domain"]
            for x in all_harvest
            if x["experiment"] == f"EXP{exp}"
        }
        rawset = unique_raw_domains(all_blocks, exp)
        eset = unique_environment_domains(all_blocks, exp)

        union = sorted(hset | eset)

        report.extend([
            f"EXP{exp}",
            "-" * 118,
            f"Raw portable domains ({len(rawset)}): {sorted(rawset)}",
            f"Corrected environmental domains ({len(eset)}): {sorted(eset)}",
            f"Harvest domains ({len(hset)}): {sorted(hset)}",
            f"Corrected exact set agreement: {'YES' if eset == hset and bool(eset) else 'NO'}",
            f"Environmental-only: {sorted(eset - hset)}",
            f"Harvest-only: {sorted(hset - eset)}",
        ])

        for d in union:
            matched = d in eset and d in hset
            linkage_rows.append({
                "experiment": f"EXP{exp}",
                "domain": d,
                "in_corrected_environment": d in eset,
                "in_harvest": d in hset,
                "identity_match": matched,
                "exp1_role_swap_rule_active": exp == 1,
            })
            if matched:
                total_match += 1
            else:
                total_mismatch += 1

        if exp == 1:
            report.append("EXP1 RAW → CORRECTED DOMAIN MAP:")
            pairs = sorted({
                (b["raw_portable_domain"], b["corrected_domain"])
                for b in all_blocks
                if b["experiment"] == "EXP1"
            })
            for raw_d, corr_d in pairs:
                report.append(f"  {raw_d} -> {corr_d}")

        report.append("")

    linkage_df = pd.DataFrame(linkage_rows)
    linkage_df.to_csv(LINKAGE_CSV, index=False)

    # Validation conditions.
    expected_block_count = 36  # 12 blocks × 3 experiments
    block_count_ok = len(all_blocks) == expected_block_count

    expected_harvest_domains = 18
    harvest_count = sum(
        len({
            x["domain"]
            for x in all_harvest
            if x["experiment"] == f"EXP{exp}"
        })
        for exp in (1, 2, 3)
    )
    harvest_count_ok = harvest_count == expected_harvest_domains

    per_exp_exact = {}
    per_exp_env_counts = {}
    per_exp_harvest_counts = {}

    for exp in (1, 2, 3):
        hset = {
            x["domain"]
            for x in all_harvest
            if x["experiment"] == f"EXP{exp}"
        }
        eset = unique_environment_domains(all_blocks, exp)
        per_exp_exact[f"EXP{exp}"] = bool(eset) and eset == hset
        per_exp_env_counts[f"EXP{exp}"] = len(eset)
        per_exp_harvest_counts[f"EXP{exp}"] = len(hset)

    all_exact = all(per_exp_exact.values())
    all_six = all(
        per_exp_env_counts[f"EXP{exp}"] == 6
        and per_exp_harvest_counts[f"EXP{exp}"] == 6
        for exp in (1, 2, 3)
    )

    # Verify correction scope: correction must appear only in EXP1.
    corrected_exp1 = [
        b for b in all_blocks
        if b["experiment"] == "EXP1" and b["correction_applied"]
    ]
    corrected_other = [
        b for b in all_blocks
        if b["experiment"] != "EXP1" and b["correction_applied"]
    ]
    correction_scope_ok = len(corrected_exp1) == 12 and len(corrected_other) == 0

    structural_pass = (
        not errors
        and digest == EXPECTED_SHA256
        and block_count_ok
        and harvest_count_ok
        and all_exact
        and all_six
        and correction_scope_ok
        and total_mismatch == 0
        and total_match == 18
    )

    verdict = "PASS" if structural_pass else "REVIEW REQUIRED"

    report.extend([
        "VALIDATION SUMMARY",
        "=" * 118,
        f"Processing errors: {len(errors)}",
        f"Archive hash verified: {'YES' if digest == EXPECTED_SHA256 else 'NO'}",
        f"Portable blocks recovered: {len(all_blocks)} / expected {expected_block_count}",
        f"Harvest domains recovered: {harvest_count} / expected {expected_harvest_domains}",
        f"Identity matches: {total_match} / 18",
        f"Identity mismatches: {total_mismatch}",
        f"EXP1 exact corrected agreement: {'YES' if per_exp_exact.get('EXP1') else 'NO'}",
        f"EXP2 exact agreement: {'YES' if per_exp_exact.get('EXP2') else 'NO'}",
        f"EXP3 exact agreement: {'YES' if per_exp_exact.get('EXP3') else 'NO'}",
        f"Correction scope valid (EXP1 only): {'YES' if correction_scope_ok else 'NO'}",
    ])
    for e in errors:
        report.append(f"ERROR: {e}")

    report.extend([
        "",
        "FINAL VERDICT",
        "=" * 118,
        verdict,
    ])

    if verdict == "PASS":
        report.extend([
            "The documented EXP1 semantic role swap resolves the prior cross-table identity mismatch.",
            "All 18 corrected environmental domain identities now agree with all 18 explicit harvest domain identities.",
            "This PASS validates domain-key linkage structure; it does NOT by itself freeze the Phase-7 ML design.",
        ])
    else:
        report.append(
            "At least one structural/linkage validation condition remains unresolved."
        )

    report.extend([
        "",
        "RESEARCH LOCK STATUS",
        "-" * 118,
        f"EXP1 analytical role-swap correction documented: {'YES' if correction_scope_ok else 'NO'}",
        f"Cross-table identity-level linkage structurally validated: {'YES' if structural_pass else 'NO'}",
        "Raw workbook modified: NO",
        "Harvest timing frozen: NO",
        "Primary target frozen: NO",
        "Feature set frozen: NO",
        "Domain definition for ML evaluation frozen: NO",
        "Shift family frozen: NO",
        "Shift severity frozen: NO",
        "Augmentation frozen: NO",
        "Model configuration frozen: NO",
        "Model training performed: NO",
    ])

    REPORT.write_text("\n".join(report) + "\n", encoding="utf-8")

    manifest = {
        "phase": "P7.6c-v2",
        "purpose": "corrected identity-level environmental/harvest domain linkage",
        "p76d_header_normalization": "strip punctuation/case/spacing before matching pH",
        "archive": str(ARCHIVE),
        "archive_sha256": digest,
        "expected_sha256_match": digest == EXPECTED_SHA256,
        "correction_id": CORRECTION_ID,
        "correction_rule": CORRECTION_RULE,
        "correction_scope": "EXP1 only",
        "raw_workbook_modified": False,
        "portable_blocks_recovered": len(all_blocks),
        "harvest_domains_recovered": harvest_count,
        "identity_matches": total_match,
        "identity_mismatches": total_mismatch,
        "per_experiment_exact_agreement": per_exp_exact,
        "correction_scope_valid": correction_scope_ok,
        "verdict": verdict,
        "processing_errors": errors,
        "phase7_ml_design_frozen": False,
        "model_training_performed": False,
        "outputs": {
            "report": str(REPORT),
            "corrected_portable_blocks": str(BLOCKS_CSV),
            "domain_linkage": str(LINKAGE_CSV),
            "harvest_domain_keys": str(HARVEST_CSV),
        },
    }
    MANIFEST.write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )

    print("\n".join(report))
    print("")
    print(f"Report: {REPORT}")
    print("RAW WORKBOOKS WERE NOT MODIFIED.")
    print("NO MODEL TRAINING WAS PERFORMED.")

    return 0 if structural_pass else 1


if __name__ == "__main__":
    raise SystemExit(main())
